"""Image editing for prose books, working directly on the .docx.

Trivia books keep each illustration as a loose file on disk keyed by chapter
number (see trivia/image_edit.py). Prose books don't: their images are embedded
*inside* the document as parts referenced by `r:embed`, already sized and
positioned. So editing here swaps the bytes of that part in place, which
changes the picture without disturbing layout, sizing or paragraph structure.

Images are addressed by paragraph index — the same index book_editor.read_blocks
reports for an image block, so the editor UI can point at one directly.

Two kinds of edit, mirroring the trivia editor:
  * AI edits  — OpenAI images.edit, optionally masked so only a painted region
    changes.
  * Local edits — deterministic Pillow operations (rotate, flip, brightness,
    contrast, saturation, sharpness) that are instant and cost nothing.

Every edit snapshots the previous bytes so any change can be undone.
"""

from __future__ import annotations

import base64
import io
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document

EMBED_ATTR = (
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
)

MAX_HISTORY = 12

# Sizes the images.edit endpoint accepts.
VALID_EDIT_SIZES = {"1024x1024", "1536x1024", "1024x1536", "auto"}

ADJUST_LIMITS = {
    "brightness": (0.2, 2.5),
    "contrast": (0.2, 2.5),
    "saturation": (0.0, 3.0),
    "sharpness": (0.0, 4.0),
}

# Interiors print black and white and must not carry stray text, so the same
# constraints applied at generation time are re-appended to every edit prompt.
# An operator instruction cannot opt out of them.
PRINT_CONSTRAINTS = (
    " Keep the result in strictly neutral grayscale: only deep black and a few "
    "distinct shades of cool gray on a pure white background. No color of any "
    "kind. Absolutely no text, letters, numbers, words, labels, captions, "
    "signatures, watermarks, frames, or borders anywhere in the image."
)


class DocxImageError(Exception):
    """A problem the user can act on, surfaced straight to the editor UI."""


@dataclass
class ImageEditResult:
    para_index: int
    history_count: int
    operation: str

    def to_dict(self) -> dict[str, Any]:
        return {
            "para_index": self.para_index,
            "history_count": self.history_count,
            "operation": self.operation,
        }


# ---------------------------------------------------------------------------
# Locating an image inside the document
# ---------------------------------------------------------------------------

def _image_part(doc, para_index: int):
    """Return the image part embedded at `para_index`.

    Raises DocxImageError rather than returning None so every caller reports
    the same user-facing message.
    """
    paras = doc.paragraphs
    if para_index < 0 or para_index >= len(paras):
        raise DocxImageError("That paragraph is not in the document.")

    blips = paras[para_index]._p.xpath(".//a:blip")
    if not blips:
        raise DocxImageError("There is no image in that paragraph.")

    embed = blips[0].get(EMBED_ATTR)
    if not embed:
        raise DocxImageError("That image is linked, not embedded, so it can't be edited.")

    try:
        return doc.part.related_parts[embed]
    except KeyError as exc:
        raise DocxImageError("The image data is missing from the document.") from exc


def list_images(path: Path) -> list[dict[str, Any]]:
    """Every embedded image, with the paragraph index that anchors it.

    `heading` is the nearest preceding heading, so the UI can name each image
    without the caller re-walking the document.
    """
    import openclaw_docx_writer as writer

    doc = Document(str(path))
    out: list[dict[str, Any]] = []
    current_heading = ""

    for i, p in enumerate(doc.paragraphs):
        if writer.is_heading_paragraph(p) or writer.is_subheading_paragraph(p):
            text = (p.text or "").strip()
            if text:
                current_heading = text
        blips = p._p.xpath(".//a:blip")
        if not blips:
            continue
        embed = blips[0].get(EMBED_ATTR)
        part = doc.part.related_parts.get(embed) if embed else None
        entry: dict[str, Any] = {
            "para_index": i,
            "heading": current_heading,
            "content_type": getattr(part, "content_type", "") if part else "",
            "bytes": len(part.blob) if part else 0,
        }
        if part is not None:
            try:
                from PIL import Image
                with Image.open(io.BytesIO(part.blob)) as im:
                    entry["width"], entry["height"] = im.size
            except Exception:
                # Dimensions are cosmetic; never fail the listing over them.
                pass
        out.append(entry)
    return out


def image_info(path: Path, para_index: int) -> dict[str, Any]:
    """Dimensions, DPI and undo depth for one image."""
    from PIL import Image

    doc = Document(str(path))
    try:
        part = _image_part(doc, para_index)
    except DocxImageError:
        return {"has_image": False}

    blob = part.blob
    try:
        with Image.open(io.BytesIO(blob)) as img:
            w, h = img.size
            dpi = img.info.get("dpi", (72, 72))
    except Exception:
        w = h = 0
        dpi = (72, 72)

    return {
        "has_image": True,
        "width": w,
        "height": h,
        "dpi": int(dpi[0]) if dpi else 72,
        "bytes": len(blob),
        "history_count": len(history_list(path, para_index)),
    }


# ---------------------------------------------------------------------------
# History
# ---------------------------------------------------------------------------

def _history_dir(docx_path: Path, para_index: int) -> Path:
    d = docx_path.parent / ".image_history" / f"{docx_path.stem}_p{para_index}"
    d.mkdir(parents=True, exist_ok=True)
    return d


def history_list(docx_path: Path, para_index: int) -> list[Path]:
    d = _history_dir(docx_path, para_index)
    return sorted((p for p in d.iterdir() if p.suffix == ".bin"), reverse=True)


def _snapshot(docx_path: Path, para_index: int, blob: bytes) -> None:
    """Store the pre-edit bytes so the change can be undone."""
    d = _history_dir(docx_path, para_index)
    stamp = f"{time.time():.6f}".replace(".", "")
    (d / f"{stamp}.bin").write_bytes(blob)

    for old in history_list(docx_path, para_index)[MAX_HISTORY:]:
        old.unlink(missing_ok=True)


# ---------------------------------------------------------------------------
# Writing an image back
# ---------------------------------------------------------------------------

def _write_image(
    docx_path: Path,
    para_index: int,
    new_bytes: bytes,
    operation: str,
    *,
    snapshot: bool = True,
) -> ImageEditResult:
    """Swap the bytes of the image part at `para_index` and save the document."""
    doc = Document(str(docx_path))
    part = _image_part(doc, para_index)

    if snapshot:
        _snapshot(docx_path, para_index, part.blob)

    # Replacing the part's blob keeps the existing relationship, so the
    # drawing's size and anchoring are untouched.
    part._blob = new_bytes
    doc.save(str(docx_path))

    return ImageEditResult(
        para_index=para_index,
        history_count=len(history_list(docx_path, para_index)),
        operation=operation,
    )


def _finalize(data: bytes, *, grayscale: bool = True) -> bytes:
    """Force neutral gray and tag print DPI, as at generation time.

    Image models routinely return a slight colour cast even when asked for
    grayscale, and these interiors print black and white — so convert rather
    than trust.
    """
    from PIL import Image

    try:
        with Image.open(io.BytesIO(data)) as src:
            img = src.convert("L").convert("RGB") if grayscale else src.convert("RGB")
            buf = io.BytesIO()
            img.save(buf, format="PNG", dpi=(300, 300))
            return buf.getvalue()
    except Exception:
        # A conversion failure must not lose the edit itself.
        return data


def current_bytes(docx_path: Path, para_index: int) -> tuple[bytes, str]:
    doc = Document(str(docx_path))
    part = _image_part(doc, para_index)
    return part.blob, getattr(part, "content_type", "image/png") or "image/png"


# ---------------------------------------------------------------------------
# Local (Pillow) edits
# ---------------------------------------------------------------------------

def adjust_image(
    docx_path: Path,
    para_index: int,
    *,
    brightness: float = 1.0,
    contrast: float = 1.0,
    saturation: float = 1.0,
    sharpness: float = 1.0,
) -> ImageEditResult:
    from PIL import Image, ImageEnhance

    for name, value in (
        ("brightness", brightness), ("contrast", contrast),
        ("saturation", saturation), ("sharpness", sharpness),
    ):
        lo, hi = ADJUST_LIMITS[name]
        if not (lo <= value <= hi):
            raise DocxImageError(f"{name} must be between {lo} and {hi}.")

    blob, _ = current_bytes(docx_path, para_index)
    with Image.open(io.BytesIO(blob)) as src:
        img = src.convert("RGB")
    for enhancer_cls, value in (
        (ImageEnhance.Brightness, brightness),
        (ImageEnhance.Contrast, contrast),
        (ImageEnhance.Color, saturation),
        (ImageEnhance.Sharpness, sharpness),
    ):
        if abs(value - 1.0) > 1e-6:
            img = enhancer_cls(img).enhance(value)

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return _write_image(docx_path, para_index, _finalize(buf.getvalue()), "adjust")


def transform_image(
    docx_path: Path,
    para_index: int,
    *,
    op: str,
    degrees: int = 90,
    axis: str = "horizontal",
) -> ImageEditResult:
    from PIL import Image, ImageOps

    blob, _ = current_bytes(docx_path, para_index)
    with Image.open(io.BytesIO(blob)) as src:
        img = src.convert("RGB")

        if op == "rotate":
            if degrees not in (90, 180, 270, -90):
                raise DocxImageError("Rotation must be 90, 180 or 270 degrees.")
            img = img.rotate(-degrees, expand=True)
        elif op == "flip":
            if axis not in {"horizontal", "vertical"}:
                raise DocxImageError("axis must be 'horizontal' or 'vertical'.")
            img = ImageOps.mirror(img) if axis == "horizontal" else ImageOps.flip(img)
        else:
            raise DocxImageError("op must be 'rotate' or 'flip'.")

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return _write_image(docx_path, para_index, _finalize(buf.getvalue()), op)


def undo(docx_path: Path, para_index: int) -> ImageEditResult:
    """Step back one edit, restoring the most recent snapshot.

    A plain stack: pop the newest snapshot and restore it. Parking the current
    bytes as a redo point would push a *newer* entry onto the same timestamp-
    ordered stack, so the next undo would restore that instead of stepping
    further back — repeated undo would oscillate between two states.
    """
    snaps = history_list(docx_path, para_index)
    if not snaps:
        raise DocxImageError("There is nothing to undo for this image.")

    latest = snaps[0]
    restored = latest.read_bytes()
    latest.unlink(missing_ok=True)

    # snapshot=False: restoring is not itself an edit worth recording.
    return _write_image(docx_path, para_index, restored, "undo", snapshot=False)


def replace_with_upload(docx_path: Path, para_index: int, data: bytes) -> ImageEditResult:
    """Swap in a user-supplied image file."""
    from PIL import Image

    try:
        with Image.open(io.BytesIO(data)) as im:
            im.verify()
    except Exception as exc:
        raise DocxImageError(f"That file is not a readable image: {exc}") from exc

    return _write_image(docx_path, para_index, _finalize(data), "upload")


# ---------------------------------------------------------------------------
# AI edits
# ---------------------------------------------------------------------------

def _decode_mask(mask_data_url: str, size: tuple[int, int]) -> bytes:
    """Turn a client-drawn mask into the RGBA PNG the API expects.

    OpenAI regenerates *transparent* pixels. The client paints the area to
    change in opaque colour on a transparent canvas, so alpha is inverted here.
    """
    from PIL import Image

    if "," in mask_data_url:
        mask_data_url = mask_data_url.split(",", 1)[1]
    try:
        raw = base64.b64decode(mask_data_url)
    except Exception as exc:
        raise DocxImageError(f"Could not decode the mask image: {exc}") from exc

    with Image.open(io.BytesIO(raw)) as m:
        mask = m.convert("RGBA")
    if mask.size != size:
        mask = mask.resize(size, Image.LANCZOS)

    alpha = mask.getchannel("A").point(lambda a: 0 if a > 8 else 255)
    out = Image.new("RGBA", size, (0, 0, 0, 255))
    out.putalpha(alpha)

    buf = io.BytesIO()
    out.save(buf, format="PNG")
    return buf.getvalue()


def ai_edit_image(
    docx_path: Path,
    para_index: int,
    instruction: str,
    *,
    mask_data_url: str = "",
    size: str = "auto",
    quality: str = "high",
    api_key: str = "",
) -> ImageEditResult:
    """Edit an embedded illustration with a natural-language instruction.

    With a mask, only the painted region changes; without one, the whole image
    is reinterpreted against the instruction.
    """
    from PIL import Image

    instruction = (instruction or "").strip()
    if not instruction:
        raise DocxImageError("Describe the change you want to make.")
    if size not in VALID_EDIT_SIZES:
        raise DocxImageError(f"size must be one of {sorted(VALID_EDIT_SIZES)}.")

    import openclaw_image_maker as image_maker

    key = image_maker.resolve_api_key(api_key)
    if not key:
        raise DocxImageError("No OpenAI API key available for image editing.")

    blob, _ = current_bytes(docx_path, para_index)

    # images.edit requires PNG input; convert whatever the document holds.
    with Image.open(io.BytesIO(blob)) as src:
        rgba = src.convert("RGBA")
        dimensions = rgba.size
        png_buf = io.BytesIO()
        rgba.save(png_buf, format="PNG")
    png_bytes = png_buf.getvalue()

    from openai import OpenAI
    client = OpenAI(api_key=key)

    kwargs: dict[str, Any] = {
        "model": "gpt-image-1",
        "image": ("illustration.png", png_bytes, "image/png"),
        "prompt": instruction + PRINT_CONSTRAINTS,
        "n": 1,
        "size": size,
        "quality": quality,
    }
    if mask_data_url:
        kwargs["mask"] = (
            "mask.png", _decode_mask(mask_data_url, dimensions), "image/png",
        )

    try:
        response = client.images.edit(**kwargs)
    except Exception as exc:
        raise DocxImageError(f"Image edit failed: {exc}") from exc

    first = response.data[0] if response.data else None
    b64 = getattr(first, "b64_json", None)
    if not b64 and isinstance(first, dict):
        b64 = first.get("b64_json")
    if not b64:
        raise DocxImageError("The image service returned no image data.")

    return _write_image(
        docx_path, para_index, _finalize(base64.b64decode(b64)), "ai_edit",
    )
