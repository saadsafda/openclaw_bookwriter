"""
manuscript_audit.py

Whole-manuscript proofreading pass. Reads a finished .docx and reports the
patterns that only exist ACROSS sections — repeated opener shapes, positional
molds ("every fourth paragraph does the same thing"), pet phrases, and uniform
section/paragraph rhythm.

This is deliberately a different scope from the checks in
openclaw_docx_writer.py. Those run inside _generate_clean_paragraph_once, on a
single freshly-generated paragraph, so a paragraph that looks fine on its own
always passes — even when all thirty sections open the same way. Formulaic
writing is a property of the book, not of any one paragraph, so it can only be
seen after assembly. Hemingway (clarity_agent.py) is sentence-local for the
same reason and cannot see any of this either.

Pure Python, no API calls: safe to run over a shelf of finished books to see
what it would have caught.

Usage:
  python3 manuscript_audit.py book.docx
  python3 manuscript_audit.py book.docx --json report.json
  python3 manuscript_audit.py story_outputs/*.docx --quiet
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from docx import Document

import openclaw_docx_writer as writer
from openclaw_docx_writer import (
    find_and_compound_monotony,
    find_invented_name_opener,
    find_problem_sentences,
    find_second_person_opener,
    find_template_sentences,
    is_heading_paragraph,
    paragraph_looks_like_body,
    split_sentences,
)

# ── Tunables ────────────────────────────────────────────────────────
# An opener shape repeated this many times across the book reads as a template.
OPENER_SHAPE_MIN = 3
# A positional mold needs this many sections before "every Nth paragraph" means
# anything -- with 3 sections, two matches is coincidence.
POSITION_MIN_SECTIONS = 4
# Share of sections that must agree at a position for it to count as a mold.
POSITION_SHARE = 0.5
# Pet phrases: n-gram length and how many uses before it is worth flagging.
PET_NGRAM_SIZES = (3, 4, 5)
PET_MIN_USES = 4
# Section lengths within this ratio of each other read as machine-uniform.
UNIFORM_SPREAD = 0.15
UNIFORM_MIN_SECTIONS = 5

# Function words carry no voice, so an n-gram made only of them is not a tell.
_STOPWORD_ONLY = {
    "a", "an", "and", "as", "at", "but", "by", "for", "from", "if", "in",
    "into", "is", "it", "its", "of", "on", "or", "that", "the", "then",
    "there", "they", "this", "to", "was", "were", "with", "you", "your",
    "he", "she", "his", "her", "we", "our", "not", "be", "been", "are",
    "have", "has", "had", "do", "does", "did", "can", "could", "would",
    "will", "when", "what", "who", "how", "all", "one", "out", "up", "so",
}


# ── Structure extraction ────────────────────────────────────────────

class Section:
    """One chapter/heading and the body paragraphs beneath it."""

    def __init__(self, title: str, index: int):
        self.title = title
        self.index = index
        self.paragraphs: list[str] = []

    @property
    def word_count(self) -> int:
        return sum(len(p.split()) for p in self.paragraphs)


def read_sections(path: Path) -> list[Section]:
    """Split a .docx into sections keyed by heading, keeping only body prose.

    Reuses is_heading_paragraph/paragraph_looks_like_body so the audit sees
    exactly the same paragraphs the writer considers body text -- otherwise the
    audit would flag front matter, captions and outline stubs as prose.
    """
    doc = Document(str(path))
    sections: list[Section] = []
    current = Section("(front matter)", 0)
    for p in doc.paragraphs:
        if is_heading_paragraph(p):
            if current.paragraphs:
                sections.append(current)
            current = Section((p.text or "").strip(), len(sections))
            continue
        if paragraph_looks_like_body(p):
            text = (p.text or "").strip()
            if text:
                current.paragraphs.append(text)
    if current.paragraphs:
        sections.append(current)
    return sections


# ── Opener shape ────────────────────────────────────────────────────

_QUOTE_STRIP = re.compile(r"^[\"'“‘(\s]+")


def opener_shape(paragraph: str, known_names: set[str] | None = None) -> str:
    """Reduce a paragraph's first sentence to a coarse grammatical shape.

    The point is to collapse different *words* that share the same *move*, so
    "Sarah opened the door" and "Marcus grabbed his coat" both become
    "NAME + past-verb". Comparing raw text would never match, which is why the
    per-paragraph detectors miss the repetition that makes a book feel canned.
    """
    sentences = split_sentences(paragraph)
    if not sentences:
        return ""
    first = _QUOTE_STRIP.sub("", sentences[0])
    words = re.findall(r"[A-Za-z']+", first)
    if not words:
        return ""
    head = words[0].lower()

    if first.rstrip().endswith("?"):
        return "question opener"
    if head in {"but", "and", "yet", "so", "still", "then"}:
        return f"conjunction opener ('{head}')"
    if head in {"imagine", "picture", "consider", "suppose", "think"}:
        return f"invitation opener ('{head}')"
    if head in {"you", "your"}:
        return "second-person opener"
    if head in {"it", "this", "that", "there", "here",
                "they", "we", "he", "she", "i"}:
        return f"pronoun opener ('{head}')"
    if head in {"the", "a", "an"}:
        return f"article opener ('{head}')"
    if head in {"when", "while", "after", "before", "as", "once", "if"}:
        return f"subordinate-clause opener ('{head}')"
    if head in {"in", "on", "at", "under", "over", "across", "beyond", "inside"}:
        return f"prepositional opener ('{head}')"
    if head in {"every", "each", "most", "some", "many", "all"}:
        return f"quantifier opener ('{head}')"
    # The "invented character" move. Every sentence starts capitalised, so
    # capitalisation alone says nothing -- "Dust hung over the valley" is not a
    # character. Two independent signals are used, because each misses cases the
    # other catches: the writer's own name+action-verb detector (which knows a
    # character introduced only ever at sentence-start), and the mid-sentence
    # capitalisation lexicon (which knows a proper noun used without a verb).
    if find_invented_name_opener(first):
        return "proper-noun + verb opener"
    if known_names and words[0] in known_names:
        if len(words) > 1 and re.match(r"[a-z]+(ed|s)$", words[1]):
            return "proper-noun + verb opener"
        return "proper-noun opener"
    if re.match(r"[a-z]+(ing)$", head):
        return "gerund opener"
    return f"other ('{head}')"


# ── Cross-section detectors ─────────────────────────────────────────

def collect_proper_nouns(sections: list[Section]) -> set[str]:
    """Words that appear capitalised somewhere other than a sentence start.

    Sentence-initial capitalisation is grammatical, not semantic, so it cannot
    distinguish "Jackson checked the fuel line" from "Dust hung over the
    valley". A word that is also capitalised mid-sentence is a real proper noun.
    """
    names: set[str] = set()
    for sec in sections:
        for para in sec.paragraphs:
            for sentence in split_sentences(para):
                tokens = re.findall(r"[A-Za-z']+", sentence)
                for tok in tokens[1:]:  # skip the sentence-initial word
                    if tok[0].isupper() and tok.lower() not in _STOPWORD_ONLY:
                        names.add(tok)
    return names



def audit_opener_shapes(sections: list[Section],
                        known_names: set[str] | None = None) -> list[dict[str, Any]]:
    """Opener shapes repeated across the whole manuscript."""
    shapes: Counter[str] = Counter()
    examples: dict[str, list[str]] = defaultdict(list)
    total = 0
    for sec in sections:
        for para in sec.paragraphs:
            shape = opener_shape(para, known_names)
            if not shape:
                continue
            total += 1
            shapes[shape] += 1
            if len(examples[shape]) < 3:
                examples[shape].append(split_sentences(para)[0][:100])
    findings = []
    for shape, count in shapes.most_common():
        if count < OPENER_SHAPE_MIN:
            continue
        share = count / total if total else 0
        if share < 0.15:
            continue
        findings.append({
            "kind": "repeated-opener-shape",
            "detail": shape,
            "count": count,
            "share": round(share, 3),
            "examples": examples[shape],
        })
    return findings


def audit_positional_molds(sections: list[Section],
                           known_names: set[str] | None = None) -> list[dict[str, Any]]:
    """The 'every time in the fourth paragraph' check.

    Groups paragraphs by their ordinal position within their section and looks
    for a shape that dominates that slot. Nothing in the per-paragraph pipeline
    can see this, because each of those paragraphs is individually fine.
    """
    multi = [s for s in sections if len(s.paragraphs) >= 2]
    if len(multi) < POSITION_MIN_SECTIONS:
        return []
    by_position: dict[int, list[tuple[str, str, str]]] = defaultdict(list)
    for sec in multi:
        for pos, para in enumerate(sec.paragraphs):
            by_position[pos].append(
                (opener_shape(para, known_names), sec.title, para))

    findings = []
    for pos in sorted(by_position):
        entries = by_position[pos]
        if len(entries) < POSITION_MIN_SECTIONS:
            continue
        shapes = Counter(shape for shape, _, _ in entries if shape)
        if not shapes:
            continue
        shape, count = shapes.most_common(1)[0]
        share = count / len(entries)
        if share < POSITION_SHARE or count < POSITION_MIN_SECTIONS:
            continue
        findings.append({
            "kind": "positional-mold",
            "detail": f"paragraph {pos + 1} of each section: {shape}",
            "count": count,
            "of": len(entries),
            "share": round(share, 3),
            "examples": [
                f"[{title[:40]}] {split_sentences(para)[0][:90]}"
                for s, title, para in entries if s == shape
            ][:3],
        })
    return findings


def audit_pet_phrases(sections: list[Section]) -> list[dict[str, Any]]:
    """Content-bearing n-grams reused across the book.

    Complements enforce_chapter_just_budget, which caps one known filler word;
    this finds the phrases a given run happens to fall in love with.
    """
    text = " ".join(p for sec in sections for p in sec.paragraphs).lower()
    words = re.findall(r"[a-z\']+", text)
    if len(words) < min(PET_NGRAM_SIZES):
        return []

    # Find the MAXIMAL repeated phrase directly, by growing each seed window as
    # far right as its repetition count holds. Stitching overlapping n-grams
    # after the fact is unsound -- two windows can share an overlap without ever
    # being adjacent in the text, which fabricates a phrase that was never
    # written ("the creek rose over the creek rose over the ford").
    seed = min(PET_NGRAM_SIZES)
    positions: dict[tuple[str, ...], list[int]] = defaultdict(list)
    for i in range(len(words) - seed + 1):
        positions[tuple(words[i:i + seed])].append(i)

    spans: list[tuple[str, int]] = []
    for gram, starts in positions.items():
        if len(starts) < PET_MIN_USES:
            continue
        if all(w in _STOPWORD_ONLY for w in gram):
            continue
        # Extend right while every occurrence still agrees on the next word.
        length = seed
        active = list(starts)
        while True:
            nxt = {words[s + length] for s in active if s + length < len(words)}
            if len(nxt) != 1 or len(active) < PET_MIN_USES:
                break
            length += 1
        spans.append((" ".join(words[starts[0]:starts[0] + length]), len(starts)))

    # A seed inside a longer maximal span at the same frequency adds nothing.
    spans.sort(key=lambda s: (-len(s[0].split()), -s[1]))
    kept: list[tuple[str, int, int]] = []
    for phrase, count in spans:
        if any(phrase in k and count <= kc for k, kc, _ in kept):
            continue
        kept.append((phrase, count, len(phrase.split())))

    # A verbatim-duplicated paragraph yields a very long span; show the head and
    # the true length rather than wrapping 200 words across the terminal.
    def _label(phrase: str, n: int) -> str:
        if n <= 12:
            return f'"{phrase}"'
        head = " ".join(phrase.split()[:12])
        return f'"{head}..." ({n} words)'

    findings = [
        {"kind": "pet-phrase", "detail": _label(phrase, n),
         "phrase": phrase, "count": count, "words": n}
        for phrase, count, n in kept
    ]
    findings.sort(key=lambda f: (-f["count"], -f["words"]))
    return findings[:15]


def audit_uniformity(sections: list[Section]) -> list[dict[str, Any]]:
    """Machine-even section lengths and paragraph counts.

    Human books vary; a generator that fills one heading at a time produces
    sections that are all within a few percent of each other.
    """
    body = [s for s in sections if s.paragraphs and s.title != "(front matter)"]
    if len(body) < UNIFORM_MIN_SECTIONS:
        return []
    findings = []
    counts = [s.word_count for s in body]
    lo, hi = min(counts), max(counts)
    if hi and (hi - lo) / hi < UNIFORM_SPREAD:
        findings.append({
            "kind": "uniform-section-length",
            "detail": f"all {len(body)} sections within "
                      f"{round((hi - lo) / hi * 100)}% of each other "
                      f"({lo}-{hi} words)",
            "count": len(body),
        })
    para_counts = Counter(len(s.paragraphs) for s in body)
    common, n = para_counts.most_common(1)[0]
    if n / len(body) >= 0.8 and len(body) >= UNIFORM_MIN_SECTIONS:
        findings.append({
            "kind": "uniform-paragraph-count",
            "detail": f"{n} of {len(body)} sections have exactly "
                      f"{common} paragraphs",
            "count": n,
        })
    return findings


def audit_closers(sections: list[Section]) -> list[dict[str, Any]]:
    """Repeated closing moves -- the summarising last line every section ends on."""
    body = [s for s in sections if s.paragraphs]
    if len(body) < POSITION_MIN_SECTIONS:
        return []
    shapes: Counter[str] = Counter()
    examples: dict[str, list[str]] = defaultdict(list)
    for sec in body:
        sentences = split_sentences(sec.paragraphs[-1])
        if not sentences:
            continue
        last = sentences[-1]
        words = re.findall(r"[A-Za-z']+", last)
        if not words:
            continue
        shape = None
        if len(words) <= 8:
            shape = "short aphoristic closer (<=8 words)"
        elif re.search(r"\b(that'?s|this is|it'?s)\b.*\b(what|why|how|the point|"
                       r"the whole|all)\b", last, re.IGNORECASE):
            shape = "summarising 'that's why/what' closer"
        elif last.rstrip().endswith("?"):
            shape = "rhetorical-question closer"
        if shape:
            shapes[shape] += 1
            if len(examples[shape]) < 3:
                examples[shape].append(last[:100])
    findings = []
    for shape, count in shapes.most_common():
        if count < POSITION_MIN_SECTIONS or count / len(body) < POSITION_SHARE:
            continue
        findings.append({
            "kind": "repeated-closer",
            "detail": shape,
            "count": count,
            "of": len(body),
            "share": round(count / len(body), 3),
            "examples": examples[shape],
        })
    return findings


def audit_existing_detectors(sections: list[Section]) -> list[dict[str, Any]]:
    """Run the writer's own per-paragraph detectors over the finished book.

    They already exist but only ever run at generation time, and with retries
    disabled for cost their findings are logged and shipped anyway. Re-running
    them here shows what actually survived into the published manuscript.
    """
    text = "\n\n".join(p for sec in sections for p in sec.paragraphs)
    short, hard = find_problem_sentences(text)
    checks = [
        ("template-sentence", find_template_sentences(text)),
        ("and-compound-monotony", find_and_compound_monotony(text)),
        ("second-person-opener", find_second_person_opener(text)),
        ("invented-name-opener", find_invented_name_opener(text)),
        ("very-hard-sentence", hard),
        ("too-short-sentence", short),
    ]
    findings = []
    for kind, hits in checks:
        if not hits:
            continue
        findings.append({
            "kind": kind,
            "detail": f"{len(hits)} occurrence(s) survived into the final text",
            "count": len(hits),
            "examples": [h[:100] for h in hits[:3]],
        })
    return findings


# ── Report ──────────────────────────────────────────────────────────

def audit(path: Path) -> dict[str, Any]:
    sections = read_sections(path)
    body = [s for s in sections if s.title != "(front matter)"]
    names = collect_proper_nouns(sections)
    findings: list[dict[str, Any]] = []
    findings += audit_positional_molds(sections, names)
    findings += audit_opener_shapes(sections, names)
    findings += audit_closers(sections)
    findings += audit_uniformity(sections)
    findings += audit_pet_phrases(sections)
    findings += audit_existing_detectors(sections)
    return {
        "file": str(path),
        "sections": len(body),
        "paragraphs": sum(len(s.paragraphs) for s in sections),
        "words": sum(s.word_count for s in sections),
        "findings": findings,
    }


_HEADLINE = {
    "positional-mold": "POSITIONAL MOLD",
    "repeated-opener-shape": "REPEATED OPENER",
    "repeated-closer": "REPEATED CLOSER",
    "uniform-section-length": "UNIFORMITY",
    "uniform-paragraph-count": "UNIFORMITY",
    "pet-phrase": "PET PHRASE",
}


def print_report(report: dict[str, Any], quiet: bool = False) -> None:
    name = Path(report["file"]).name
    print(f"\n=== {name} ===")
    print(f"{report['sections']} sections · {report['paragraphs']} paragraphs · "
          f"{report['words']} words")
    if not report["findings"]:
        print("No cross-section patterns found.")
        return
    for f in report["findings"]:
        label = _HEADLINE.get(f["kind"], f["kind"].replace("-", " ").upper())
        line = f"\n  [{label}] {f['detail']}"
        if "of" in f:
            line += f"  ({f['count']}/{f['of']}, {int(f['share'] * 100)}%)"
        elif f.get("count") and "share" in f:
            line += f"  ({f['count']}×, {int(f['share'] * 100)}% of paragraphs)"
        elif f.get("count"):
            line += f"  ({f['count']}×)"
        print(line)
        if not quiet:
            for ex in f.get("examples", []):
                print(f"      · {ex}")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("paths", nargs="+", type=Path, help=".docx file(s) to audit")
    ap.add_argument("--json", type=Path, help="write full report as JSON")
    ap.add_argument("--quiet", action="store_true", help="omit example lines")
    args = ap.parse_args(argv)

    reports = []
    for path in args.paths:
        if not path.exists():
            print(f"SKIP (missing): {path}", file=sys.stderr)
            continue
        try:
            report = audit(path)
        except Exception as e:  # a bad docx must not kill a batch run
            print(f"SKIP ({type(e).__name__}: {e}): {path}", file=sys.stderr)
            continue
        reports.append(report)
        print_report(report, quiet=args.quiet)

    if args.json and reports:
        args.json.write_text(json.dumps(reports, indent=2))
        print(f"\nJSON report: {args.json}")
    return 0 if reports else 1


if __name__ == "__main__":
    raise SystemExit(main())
