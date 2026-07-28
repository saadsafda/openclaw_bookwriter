from __future__ import annotations

import json
import os
import re
import shlex
import signal
import subprocess
import sys
import threading
import time
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from docx.text.paragraph import Paragraph
from flask import Flask, abort, jsonify, render_template, request, send_file
from werkzeug.utils import secure_filename

import io
import base64
import qrcode
from qrcode.constants import ERROR_CORRECT_L, ERROR_CORRECT_M, ERROR_CORRECT_Q, ERROR_CORRECT_H
from qrcode.image.styledpil import StyledPilImage
from qrcode.image.styles.moduledrawers.pil import (
    SquareModuleDrawer,
    RoundedModuleDrawer,
    CircleModuleDrawer,
    GappedSquareModuleDrawer,
    VerticalBarsDrawer,
    HorizontalBarsDrawer,
)
from qrcode.image.styles.colormasks import SolidFillColorMask

import openclaw_image_maker as image_maker
import openclaw_docx_writer as writer
import pub_listing_agent
import wp_landing_page
import db as bookdb
import publications as pub_routes
import review_automation as review_routes
import launch_emails as launch_email_routes
import book_editor as book_editor_routes

ROOT_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = ROOT_DIR / "web_uploads"
OUTPUT_DIR = ROOT_DIR / "web_outputs"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
PYTHON_BIN = ROOT_DIR / ".venv" / "bin" / "python"
if not PYTHON_BIN.exists():
    PYTHON_BIN = Path(sys.executable)

DEFAULTS: dict[str, Any] = {
    "agent": "main",
    "tone": "friendly, encouraging, and easy to understand",
    # WHAT the book is (subject, angle, audience) and HOW it must sound (tone,
    # register), kept separate: merged into one field the model absorbs the
    # subject and drops the tone. Both are prepended to every paragraph prompt
    # so a generic-sounding heading is written to fit the book.
    "book_premise": "",
    "book_voice": "",
    # Pre-split single field, retained so saved jobs keep working.
    "book_context": "",
    # Let the writer infer premise/voice from the outline when the user leaves
    # them blank. One extra model call per book, before any paragraph is written.
    "auto_book_context": True,
    "words": 250,
    "words_max": 320,
    "subwords": 250,
    "subwords_max": 320,
    # Empty string used to mean "let OpenClaw pick" — but OpenClaw's own
    # default in that case is "adaptive" thinking (confirmed from a live
    # call's requestShaping.thinking field), which lets the model reason at
    # open-ended length before writing each paragraph. That's real added
    # latency and cost for a task (write a ~300-word paragraph from a
    # detailed, fully-specified prompt) that doesn't need deep reasoning.
    # "off" skips that reasoning step entirely.
    "thinking": "off",
    "timeout": 180,
    "images": True,
    "force": False,
    "image_prompt_variant": "rich-scene-no-text",
    "image_model": "gpt-image-1",
    "image_size": "1024x1536",
    "image_quality": "high",
    "image_width": 5.5,
    # Warning threshold for real OpenClaw spend per generation run. Reaching
    # it opens a Continue/Stop modal, but the writer keeps running unless the
    # user explicitly chooses Stop.
    "max_spend_usd": 15.0,
}

LONG_BOOK_MIN_PAGES = 600
LONG_BOOK_AGENT_ID = "long-writer-agent-1"
LONG_BOOK_MODEL_KEY = "claude-sonnet-4.6"
DASHBOARD_MODE_STANDARD = "standard"
DASHBOARD_MODE_LONG = "long-book"


@dataclass
class Job:
    id: str
    input_docx: str
    output_docx: str
    final_docx: str
    kindle_docx: str = ""
    paperback_docx: str = ""
    status: str = "queued"
    current_action: str = ""
    error: str = ""
    headings: list[str] = field(default_factory=list)
    listing: dict[str, Any] = field(default_factory=dict)
    logs: list[str] = field(default_factory=list)
    config: dict[str, Any] = field(default_factory=dict)
    pre_written: bool = False
    custom_title: str = ""
    hemingway_login_required: bool = False
    # Set when the writer crosses the configured spend warning. Generation
    # keeps running; the frontend uses this only to show Continue/Stop.
    budget_paused: bool = False
    budget_spent_usd: float = 0.0
    budget_limit_usd: float = 0.0
    stop_requested: bool = False
    active_process: subprocess.Popen[str] | None = field(default=None, repr=False)
    created_at: float = field(default_factory=time.time)
    updated_at: float = field(default_factory=time.time)
    lock: threading.Lock = field(default_factory=threading.Lock, repr=False)


@dataclass
class Batch:
    """A set of books selected together in the UI and generated one at a time.

    A batch owns no generation logic of its own: it holds an ordered list of
    normal Job ids and a worker thread that runs each job's existing
    _run_generation to completion before starting the next. Every per-book
    feature (stop, spend cap, downloads, DB persistence) keeps working through
    the underlying Job untouched.
    """
    id: str
    job_ids: list[str] = field(default_factory=list)
    current_index: int = 0  # index into job_ids of the book being generated
    status: str = "queued"  # queued | running | done
    created_at: float = field(default_factory=time.time)
    updated_at: float = field(default_factory=time.time)
    lock: threading.Lock = field(default_factory=threading.Lock, repr=False)


app = Flask(__name__)
# Re-read templates from disk on each request so HTML edits show up without a
# restart (small per-request cost; fine for this single-user local app).
app.config["TEMPLATES_AUTO_RELOAD"] = True
JOBS: dict[str, Job] = {}
JOBS_LOCK = threading.Lock()
BATCHES: dict[str, Batch] = {}
BATCHES_LOCK = threading.Lock()

bookdb.init_db()
pub_routes.register(app)
review_routes.register(app)
review_routes.start_background_tick()
launch_email_routes.register(app)
book_editor_routes.register(app)


def _timestamp() -> str:
    return time.strftime("%H:%M:%S")


def _rehydrate_job_from_db(job_id: str) -> Job | None:
    """Rebuild a Job dataclass from a persisted book record.

    Used when the user interacts with a historical book whose in-memory Job
    was lost (e.g. server restart, or job created before this session).
    """
    book = bookdb.get_book(job_id)
    if book is None:
        return None
    final_docx = book.get("final_docx") or ""
    output_docx = final_docx or book.get("input_docx") or ""
    derived_title = _derive_title(book.get("input_docx") or job_id)
    stored_title = (book.get("title") or "").strip()
    custom_title = stored_title if stored_title and stored_title != derived_title else ""
    config = dict(book.get("config") or {})
    status = book.get("status") or "success"
    return Job(
        id=job_id,
        input_docx=book.get("input_docx") or "",
        output_docx=output_docx,
        final_docx=final_docx,
        kindle_docx=book.get("kindle_docx") or "",
        paperback_docx=book.get("paperback_docx") or "",
        status=status,
        headings=list(book.get("headings") or []),
        listing=dict(book.get("listing") or {}),
        logs=list(book.get("logs") or []),
        config=config,
        pre_written=bool(book.get("pre_written") or 0),
        custom_title=custom_title,
        budget_paused=status == "budget_paused",
        budget_spent_usd=float(config.get("_budget_spent_usd", 0) or 0),
        budget_limit_usd=float(config.get("_budget_limit_usd", 0) or 0),
        created_at=float(book.get("created_at") or time.time()),
        updated_at=float(book.get("updated_at") or time.time()),
    )


def _get_job(job_id: str) -> Job:
    with JOBS_LOCK:
        job = JOBS.get(job_id)
        if job is None:
            rehydrated = _rehydrate_job_from_db(job_id)
            if rehydrated is not None:
                JOBS[job_id] = rehydrated
                job = rehydrated
    if job is None:
        abort(404, description="Job not found")
    return job


# Log markers emitted by clarity_agent.py when Hemingway's "Simplify" button
# never appears (headless browser session not logged in).
_HEMINGWAY_LOGIN_MARKERS = (
    "hemingway is not logged in",
    "simplify feature unavailable",
    "button is not appearing",
)

# Machine-parseable line openclaw_docx_writer.py prints once when its spend
# warning threshold is crossed. The writer continues running after this line.
_BUDGET_CAP_RE = re.compile(
    r"BUDGET_CAP_HIT spent_usd=([\d.]+) limit_usd=([\d.]+) calls=(\d+)"
)


def _append_log(job: Job, text: str) -> None:
    line = f"[{_timestamp()}] {text}"
    low = text.lower()
    with job.lock:
        job.logs.append(line)
        if any(marker in low for marker in _HEMINGWAY_LOGIN_MARKERS):
            job.hemingway_login_required = True
        m = _BUDGET_CAP_RE.search(text)
        if m:
            job.budget_paused = True
            job.budget_spent_usd = float(m.group(1))
            job.budget_limit_usd = float(m.group(2))
            job.config["_budget_spent_usd"] = job.budget_spent_usd
            job.config["_budget_limit_usd"] = job.budget_limit_usd
        job.updated_at = time.time()


def _set_status(job: Job, status: str, action: str = "", error: str = "") -> None:
    with job.lock:
        job.status = status
        job.current_action = action
        job.error = error
        job.updated_at = time.time()


def _derive_title(input_path: str) -> str:
    """Extract a human-readable title from the input filename."""
    name = Path(input_path).stem
    # Strip uuid prefix added by upload (8hex_filename)
    import re
    name = re.sub(r'^[0-9a-f]{8}_', '', name)
    return name.replace('_', ' ').replace('-', ' ').strip() or 'Untitled Book'


def _sync_job_to_db(job: Job) -> None:
    """Persist current job state to SQLite."""
    with job.lock:
        title = job.custom_title.strip() or _derive_title(job.input_docx)
        bookdb.save_book(
            book_id=job.id,
            title=title,
            status=job.status,
            agent=job.config.get('agent', 'main'),
            model=job.config.get('image_model', ''),
            input_docx=job.input_docx,
            final_docx=job.final_docx,
            kindle_docx=job.kindle_docx,
            paperback_docx=job.paperback_docx,
            headings=list(job.headings),
            listing=dict(job.listing) if job.listing else {},
            config=dict(job.config),
            logs=list(job.logs),
            error=job.error,
            pre_written=bool(job.pre_written),
        )
        was_success = job.status == "success"

    # Auto-create a draft publication once the writer pipeline finishes.
    # Idempotent: pub_routes.auto_create_from_book returns existing id if any.
    if was_success:
        try:
            book = bookdb.get_book(job.id)
            if book:
                pub_routes.auto_create_from_book(book)
        except Exception:
            # Never block the writer pipeline on a publication-side failure.
            pass


def _detect_pre_written(doc_path: Path, min_body_paragraphs: int = 3) -> bool:
    """Return True if the uploaded .docx already contains substantial body prose.

    Uses writer.paragraph_looks_like_body to count paragraphs that look like
    real written prose (not headings, not bare outline topics). If the file
    has at least `min_body_paragraphs` such paragraphs, it's considered
    "already written" rather than an outline-only skeleton.
    """
    try:
        doc = Document(str(doc_path))
    except Exception:
        return False
    body_count = 0
    for p in doc.paragraphs:
        try:
            if writer.paragraph_looks_like_body(p):
                body_count += 1
                if body_count >= min_body_paragraphs:
                    return True
        except Exception:
            continue
    return False


def _list_image_headings(doc_path: Path) -> list[str]:
    doc = Document(str(doc_path))
    out: list[str] = []
    seen: set[str] = set()
    i = 0
    while i < len(doc.paragraphs):
        selected = writer._select_heading_for_image(doc, i)
        if selected is None:
            i += 1
            continue
        heading_index, _para, heading = selected
        key = heading.strip().lower()
        if heading and key not in seen:
            seen.add(key)
            out.append(heading)
        i = heading_index + 1
    return out


def _list_rewritable_headings(doc_path: Path) -> list[dict[str, Any]]:
    """List every heading that actually has a body paragraph to rewrite.

    This is deliberately NOT _list_image_headings(): that one returns only
    chapter-level headings (the right unit for one image per chapter), but
    chapter titles get no intro text by default, so the rewrite menu showed
    only entries that could not be rewritten. Walk the document instead and
    report each heading with its level, plus whether a body paragraph was
    found for it.
    """
    doc = Document(str(doc_path))
    paragraphs = doc.paragraphs
    out: list[dict[str, Any]] = []
    seen: set[str] = set()
    current_chapter = ""

    for i, p in enumerate(paragraphs):
        is_h = writer.is_heading_paragraph(p)
        is_sub = (not is_h) and writer.is_subheading_paragraph(p)
        if not is_h and not is_sub:
            continue
        heading = (p.text or "").strip()
        if not heading:
            continue

        # Front matter (TOC, bonus pages, subtitle block) is not rewritable
        # prose — offering it just invites a rewrite that can't do anything.
        flat = re.sub(r"\s+", " ", heading).strip().lower()
        flat = re.sub(r"^chapter:\s*", "", flat)
        if (flat.startswith("table of contents")
                or flat.startswith("free bonus")
                or flat.startswith("— a comprehensive guide")
                or flat.startswith("- a comprehensive guide")):
            continue

        body = writer.find_body_paragraph_after(paragraphs, i)
        if is_h:
            current_chapter = heading

        key = heading.lower()
        if key in seen:
            continue
        seen.add(key)

        out.append({
            "heading": heading,
            "level": "chapter" if is_h else "section",
            "chapter": "" if is_h else current_chapter,
            "has_body": body is not None,
            "preview": ((body.text or "").strip()[:140] if body is not None else ""),
        })
    return out


def _pick_final_doc(output_doc: Path, logs: list[str]) -> Path:
    for line in reversed(logs):
        marker = "Final file:"
        if marker in line:
            candidate = Path(line.split(marker, 1)[1].strip())
            if candidate.exists():
                return candidate

    clear_candidate = output_doc.with_stem(output_doc.stem + "_formatted_clear")
    formatted_candidate = output_doc.with_stem(output_doc.stem + "_formatted")
    for candidate in (clear_candidate, formatted_candidate, output_doc):
        if candidate.exists():
            return candidate
    return output_doc


# Exit code openclaw_docx_writer.py returns specifically when it stops itself
# because --max-spend-usd was hit (see its __main__ block). Distinct from
# every other nonzero exit so the web UI can offer Continue/Stop instead of
# reporting a plain error.
_BUDGET_CAP_EXIT_CODE = 3


class BudgetPausedError(RuntimeError):
    """The writer stopped itself on purpose because it hit its spending cap
    (not a crash). Caught separately in _run_generation to flag the job for
    the Continue/Stop modal instead of marking it a hard failure."""


class GenerationStoppedError(RuntimeError):
    """The user explicitly stopped an active generation subprocess."""


def _stream_command(job: Job, cmd: list[str]) -> None:
    with job.lock:
        if job.stop_requested:
            raise GenerationStoppedError("Generation stopped by user.")

    pretty = " ".join(shlex.quote(part) for part in cmd)
    _append_log(job, f"$ {pretty}")

    proc = subprocess.Popen(
        cmd,
        cwd=str(ROOT_DIR),
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        bufsize=1,
        start_new_session=True,
    )
    with job.lock:
        job.active_process = proc

    try:
        assert proc.stdout is not None
        for raw_line in proc.stdout:
            line = raw_line.rstrip("\n")
            if line.strip():
                _append_log(job, line)
        proc.stdout.close()
        rc = proc.wait()
    finally:
        with job.lock:
            stopped_by_user = job.stop_requested
            if job.active_process is proc:
                job.active_process = None

    if stopped_by_user:
        raise GenerationStoppedError("Generation stopped by user.")
    if rc == _BUDGET_CAP_EXIT_CODE:
        raise BudgetPausedError(
            f"Writer stopped: spending cap reached (${job.budget_spent_usd:.2f} "
            f"of ${job.budget_limit_usd:.2f})."
        )
    if rc != 0:
        raise RuntimeError(f"Command failed with exit code {rc}")


def _run_kdp_formatting(job: Job, source_doc: Path) -> tuple[Path, Path]:
    kindle_doc = source_doc.with_stem(source_doc.stem + "_kindle")
    paperback_doc = source_doc.with_stem(source_doc.stem + "_paperback")
    cfg = dict(job.config)

    cmd = [
        str(PYTHON_BIN),
        str(ROOT_DIR / "kdp_docx_formatter.py"),
        str(source_doc),
        "--kindle-output",
        str(kindle_doc),
        "--paperback-output",
        str(paperback_doc),
        "--title-placeholder",
        str(cfg.get("title_placeholder", "Book Title Placeholder")),
        "--author-placeholder",
        str(cfg.get("author_placeholder", "Author Name")),
    ]

    # The original outline is the authority on which lines are subheadings.
    # Without it the formatter guesses from text shape and promotes AI-written
    # prose (colon lead-ins, generated list items) into headings.
    outline_path = Path(job.input_docx) if job.input_docx else None
    if outline_path and outline_path.exists():
        cmd.extend(["--outline", str(outline_path)])

    estimated_pages = int(cfg.get("estimated_pages", 0) or 0)
    if estimated_pages > 0:
        cmd.extend(["--estimated-pages", str(estimated_pages)])

    _append_log(job, "Formatting KDP deliverables (Kindle + Paperback)...")
    _stream_command(job, cmd)

    if not kindle_doc.exists() or not paperback_doc.exists():
        raise RuntimeError("KDP formatter completed but output files were not found.")

    return kindle_doc, paperback_doc


def _book_context_flag(job: Job, cfg: dict[str, Any]) -> list[str]:
    """Return the --book-premise-file / --book-voice-file flags, or [].

    Each is written to a file next to the job's output rather than passed
    inline: they are free-form multi-line text from the user, and long text as
    a command-line argument is fragile.

    Premise and voice stay separate all the way to the prompt — merging them
    into one blob makes the model absorb the subject and drop the tone.
    `book_context` is the pre-split field, still read so saved jobs keep
    working; it fills in as the premise.
    """
    premise = str(cfg.get("book_premise") or cfg.get("book_context") or "").strip()
    voice = str(cfg.get("book_voice") or "").strip()

    flags: list[str] = []
    for name, value in (("premise", premise), ("voice", voice)):
        if not value:
            continue
        path = OUTPUT_DIR / f"{job.id}_book_{name}.txt"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(value, encoding="utf-8")
        flags.extend([f"--book-{name}-file", str(path)])
    return flags


def _build_generation_cmd(job: Job, cfg: dict[str, Any], max_spend_usd: float) -> list[str]:
    input_doc = Path(job.input_docx)
    cmd = [
        str(PYTHON_BIN),
        str(ROOT_DIR / "openclaw_docx_writer.py"),
        str(input_doc),
        "--agent",
        str(cfg["agent"]),
        "--no-cache",
        "--image-model",
        str(cfg["image_model"]),
        "--max-spend-usd",
        str(max_spend_usd),
    ]
    # Already-written book: never generate TEXT (--no-text). This is the fix
    # for a human-written book whose paragraph formatting confused the heading
    # detector into "filling in" content under real paragraphs (a 29k-word
    # book ballooned to 100k). --no-text short-circuits all text generation.
    if job.pre_written:
        cmd.append("--no-text")
    # Images are independent of text: when the "Generate Images" toggle is on,
    # pass --images for BOTH already-written and outline books. For an
    # already-written book the writer still runs the image pass, and
    # insert_images_into_document only creates images for headings that don't
    # already have one (existing images are left alone). Default True keeps
    # backward behavior for outline books that predate this config field.
    if cfg.get("images", True):
        cmd.append("--images")
    # Was never actually passed before — openclaw_docx_writer.py only adds
    # --thinking to the openclaw call when this is non-empty (see
    # run_openclaw_call), so an empty/missing value here silently left every
    # call on OpenClaw's own default, which is "adaptive" reasoning — real
    # added latency and cost per paragraph. Always pass it explicitly now.
    thinking = str(cfg.get("thinking") or DEFAULTS["thinking"])
    cmd.extend(["--thinking", thinking])
    ctx_flag = _book_context_flag(job, cfg)
    if ctx_flag:
        cmd.extend(ctx_flag)
    # Auto-detect the book's premise/voice from the outline unless the user
    # turned it off. The writer skips detection anyway when both were typed.
    if not cfg.get("auto_book_context", True):
        cmd.append("--no-auto-context")
    if cfg.get("openai_api_key"):
        cmd.extend(["--openai-api-key", str(cfg["openai_api_key"])])
    return cmd


def _finish_generation(job: Job, output_doc: Path) -> None:
    """Shared tail of a successful (or resumed-and-now-successful) generation
    run: pick the final doc, run KDP formatting, record results.

    For an already-written book (job.pre_written) the chosen behavior is
    formatting only — the writer already applied book formatting, so skip the
    KDP Kindle/paperback conversion here and keep the author's document as-is.
    """
    with job.lock:
        final_doc = _pick_final_doc(output_doc, job.logs)
        job.final_docx = str(final_doc)
        pre_written = bool(job.pre_written)

    headings = _list_image_headings(final_doc) if final_doc.exists() else []

    if pre_written:
        with job.lock:
            job.headings = headings
            images_on = bool(job.config.get("images", True))
        _append_log(job, f"Ready. Formatted document: {final_doc}")
        extra = " Images were generated for headings that lacked one." if images_on else ""
        _append_log(job, "Already-written book: skipped text generation and KDP "
                         "conversion; formatting was applied." + extra)
    else:
        kindle_doc, paperback_doc = _run_kdp_formatting(job, final_doc)
        with job.lock:
            job.headings = headings
            job.kindle_docx = str(kindle_doc)
            job.paperback_docx = str(paperback_doc)
        _append_log(job, f"Ready. Final document: {final_doc}")
        _append_log(job, f"Kindle output: {kindle_doc}")
        _append_log(job, f"Paperback output: {paperback_doc}")

    with job.lock:
        job.budget_paused = False
    _set_status(job, "success", action="", error="")
    _sync_job_to_db(job)


def _run_generation(job_id: str, max_spend_usd: float | None = None) -> None:
    """Run (or resume) book generation. Sections already written to the
    output .docx are skipped by the writer regardless of --no-cache (see
    skip_text_generation in openclaw_docx_writer.py), so calling this again
    after a budget pause only pays for what's still missing."""
    job = _get_job(job_id)
    cfg = dict(job.config)

    with job.lock:
        job.budget_paused = False
        job.budget_spent_usd = 0.0
        job.budget_limit_usd = 0.0
        job.stop_requested = False
    _set_status(job, "running", action="generating_book", error="")

    input_doc = Path(job.input_docx)
    output_doc = input_doc

    spend_cap = max_spend_usd if max_spend_usd is not None else float(
        cfg.get("max_spend_usd", DEFAULTS["max_spend_usd"])
    )
    cmd = _build_generation_cmd(job, cfg, spend_cap)

    try:
        _stream_command(job, cmd)
        _finish_generation(job, output_doc)
        # Landing page + QR is now manual-only (use the QR button in the UI).
    except GenerationStoppedError:
        _append_log(job, "User chose Stop. Keeping the book as generated so far.")
        with job.lock:
            job.stop_requested = False
            job.budget_paused = False
        try:
            _finish_generation(job, output_doc)
        except Exception as exc:
            _append_log(job, f"ERROR finalizing partial book: {exc}")
            _set_status(job, "error", action="", error=str(exc))
            _sync_job_to_db(job)
    except BudgetPausedError as exc:
        # Compatibility with jobs created by the older writer, which exited
        # at the threshold. New runs only warn and continue.
        _append_log(job, str(exc))
        with job.lock:
            job.config["_budget_spent_usd"] = job.budget_spent_usd
            job.config["_budget_limit_usd"] = job.budget_limit_usd
        _set_status(job, "budget_paused", action="", error="")
        _sync_job_to_db(job)
    except Exception as exc:
        _append_log(job, f"ERROR: {exc}")
        _set_status(job, "error", action="", error=str(exc))
        _sync_job_to_db(job)


def _resume_generation(job_id: str, new_max_spend_usd: float) -> None:
    """Resume a legacy job that was paused by the older hard-cap behavior."""
    _run_generation(job_id, max_spend_usd=new_max_spend_usd)


def _finalize_legacy_stopped_job(job_id: str) -> None:
    """Finalize output from a job paused by the previous hard-cap behavior."""
    job = _get_job(job_id)
    output_doc = Path(job.output_docx)
    _set_status(job, "running", action="finalizing_partial_book", error="")
    try:
        _finish_generation(job, output_doc)
    except Exception as exc:
        _append_log(job, f"ERROR finalizing partial book: {exc}")
        _set_status(job, "error", action="", error=str(exc))
        _sync_job_to_db(job)


def _run_replace_images(job_id: str, headings: list[str], overrides: dict[str, Any]) -> None:
    job = _get_job(job_id)

    with job.lock:
        final_doc = Path(job.final_docx or job.output_docx)
        base_cfg = dict(job.config)

    cfg = {**base_cfg, **overrides}

    _set_status(job, "running", action="replacing_images", error="")
    _append_log(job, f"Replacing images for {len(headings)} heading(s)...")

    try:
        for heading in headings:
            cmd = [
                str(PYTHON_BIN),
                str(ROOT_DIR / "openclaw_docx_writer.py"),
                str(final_doc),
                "--images",
                "--image-heading",
                heading,
                "--image-prompt-variant",
                str(cfg["image_prompt_variant"]),
                "--image-model",
                str(cfg["image_model"]),
                "--image-size",
                str(cfg["image_size"]),
                "--image-quality",
                str(cfg["image_quality"]),
                "--image-width",
                str(cfg["image_width"]),
            ]
            if cfg.get("openai_api_key"):
                cmd.extend(["--openai-api-key", str(cfg["openai_api_key"])])
            if cfg.get("image_guidance"):
                cmd.extend(["--image-guidance", str(cfg["image_guidance"])])

            _append_log(job, f"Heading: {heading}")
            _stream_command(job, cmd)

        headings_after = _list_image_headings(final_doc) if final_doc.exists() else []
        kindle_doc, paperback_doc = _run_kdp_formatting(job, final_doc)
        with job.lock:
            job.headings = headings_after
            job.final_docx = str(final_doc)
            job.kindle_docx = str(kindle_doc)
            job.paperback_docx = str(paperback_doc)

        _append_log(job, "Image replacement completed.")
        _append_log(job, f"Kindle output refreshed: {kindle_doc}")
        _append_log(job, f"Paperback output refreshed: {paperback_doc}")
        _set_status(job, "success", action="", error="")
        _sync_job_to_db(job)
    except Exception as exc:
        _append_log(job, f"ERROR: {exc}")
        _set_status(job, "error", action="", error=str(exc))
        _sync_job_to_db(job)


def _run_rewrite_paragraphs(job_id: str, headings: list[str], overrides: dict[str, Any]) -> None:
    job = _get_job(job_id)

    with job.lock:
        final_doc = Path(job.final_docx or job.output_docx)
        base_cfg = dict(job.config)

    cfg = {**base_cfg, **overrides}

    _set_status(job, "running", action="rewriting_paragraphs", error="")
    _append_log(job, f"Rewriting paragraphs for {len(headings)} heading(s)...")

    try:
        for heading in headings:
            cmd = [
                str(PYTHON_BIN),
                str(ROOT_DIR / "openclaw_docx_writer.py"),
                str(final_doc),
                "--agent",
                str(cfg["agent"]),
                "--rewrite-heading",
                heading,
                "--tone",
                str(cfg.get("tone", DEFAULTS["tone"])),
                "--words",
                str(cfg.get("words", DEFAULTS["words"])),
                "--words-max",
                str(cfg.get("words_max", DEFAULTS["words_max"])),
                "--subwords",
                str(cfg.get("subwords", DEFAULTS["subwords"])),
                "--subwords-max",
                str(cfg.get("subwords_max", DEFAULTS["subwords_max"])),
            ]
            ctx_flag = _book_context_flag(job, cfg)
            if ctx_flag:
                cmd.extend(ctx_flag)
            if cfg.get("rewrite_guidance"):
                cmd.extend(["--rewrite-guidance", str(cfg["rewrite_guidance"])])

            _append_log(job, f"Heading: {heading}")
            _stream_command(job, cmd)

        kindle_doc, paperback_doc = _run_kdp_formatting(job, final_doc)
        with job.lock:
            job.final_docx = str(final_doc)
            job.kindle_docx = str(kindle_doc)
            job.paperback_docx = str(paperback_doc)

        _append_log(job, "Paragraph rewrite completed.")
        _append_log(job, f"Kindle output refreshed: {kindle_doc}")
        _append_log(job, f"Paperback output refreshed: {paperback_doc}")
        _set_status(job, "success", action="", error="")
        _sync_job_to_db(job)
    except Exception as exc:
        _append_log(job, f"ERROR: {exc}")
        _set_status(job, "error", action="", error=str(exc))
        _sync_job_to_db(job)


def _bool_from_form(value: str | None) -> bool:
    if value is None:
        return False
    return value.lower() in {"1", "true", "yes", "on"}


def _normalize_output_name(name: str, input_path: Path, job_id: str) -> str:
    cleaned = secure_filename(name.strip()) if name.strip() else ""
    if not cleaned:
        cleaned = f"{input_path.stem}_generated_{job_id[:8]}.docx"
    if not cleaned.lower().endswith(".docx"):
        cleaned += ".docx"
    return cleaned


def _parse_int(value: str | None, fallback: int) -> int:
    try:
        return int(value) if value not in (None, "") else fallback
    except Exception:
        return fallback


def _parse_float(value: str | None, fallback: float) -> float:
    try:
        return float(value) if value not in (None, "") else fallback
    except Exception:
        return fallback


def _normalize_dashboard_mode(value: str | None, *, default: str = DASHBOARD_MODE_STANDARD) -> str:
    mode = (value or default).strip().lower()
    if mode not in {DASHBOARD_MODE_STANDARD, DASHBOARD_MODE_LONG}:
        return default
    return mode


def _set_openclaw_default_model(model_key: str) -> None:
    proc = subprocess.run(
        ["openclaw", "models", "set", model_key],
        capture_output=True,
        text=True,
        timeout=60,
    )
    if proc.returncode != 0:
        raise RuntimeError(proc.stderr.strip() or proc.stdout.strip() or "Unknown error")

    # Invalidate model cache when available.
    if "_MODEL_CACHE" in globals() and isinstance(_MODEL_CACHE, dict):
        _MODEL_CACHE["data"] = None
        _MODEL_CACHE["ts"] = 0.0


@app.get("/")
def index() -> str:
    return _render_writer_dashboard(is_long_book=False)


@app.get("/long-book")
def long_book_index() -> str:
    return _render_writer_dashboard(is_long_book=True)


def _render_writer_dashboard(*, is_long_book: bool) -> str:
    prompt_variants = sorted(image_maker.PROMPT_VARIANTS.keys())
    dashboard_mode = DASHBOARD_MODE_LONG if is_long_book else DASHBOARD_MODE_STANDARD
    return render_template(
        "index.html",
        defaults=DEFAULTS,
        prompt_variants=prompt_variants,
        dashboard_mode=dashboard_mode,
        is_long_book=is_long_book,
        long_book_min_pages=LONG_BOOK_MIN_PAGES,
        long_book_agent_id=LONG_BOOK_AGENT_ID,
        long_book_model_key=LONG_BOOK_MODEL_KEY,
        estimated_pages_default=(LONG_BOOK_MIN_PAGES if is_long_book else 0),
    )


@app.get("/qr-code")
def qr_code_page() -> str:
    return render_template("qr_code.html")


@app.get("/settings")
def settings_page() -> str:
    return render_template("settings.html")


_QR_EC_LEVELS = {
    "L": ERROR_CORRECT_L,
    "M": ERROR_CORRECT_M,
    "Q": ERROR_CORRECT_Q,
    "H": ERROR_CORRECT_H,
}

_QR_DRAWERS = {
    "square": SquareModuleDrawer,
    "rounded": RoundedModuleDrawer,
    "circle": CircleModuleDrawer,
    "gapped": GappedSquareModuleDrawer,
    "vertical": VerticalBarsDrawer,
    "horizontal": HorizontalBarsDrawer,
}


def _hex_to_rgb(value: str, fallback: tuple[int, int, int]) -> tuple[int, int, int]:
    if not value:
        return fallback
    v = value.strip().lstrip("#")
    if len(v) == 3:
        v = "".join(c * 2 for c in v)
    if len(v) != 6:
        return fallback
    try:
        return (int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16))
    except ValueError:
        return fallback


@app.post("/api/qr-code")
def generate_qr_code() -> Any:
    data = request.get_json(silent=True) or {}
    content = (data.get("content") or "").strip()
    if not content:
        return jsonify({"error": "content is required"}), 400
    if len(content) > 2000:
        return jsonify({"error": "content too long (max 2000 chars)"}), 400

    ec_key = (data.get("error_correction") or "M").upper()
    ec_level = _QR_EC_LEVELS.get(ec_key, ERROR_CORRECT_M)

    try:
        box_size = max(4, min(40, int(data.get("box_size") or 12)))
    except (TypeError, ValueError):
        box_size = 12
    try:
        border = max(0, min(16, int(data.get("border") or 4)))
    except (TypeError, ValueError):
        border = 4

    fg = _hex_to_rgb(data.get("fg_color") or "#111827", (17, 24, 39))
    bg = _hex_to_rgb(data.get("bg_color") or "#ffffff", (255, 255, 255))

    style_key = (data.get("style") or "square").lower()
    drawer_cls = _QR_DRAWERS.get(style_key, SquareModuleDrawer)

    qr = qrcode.QRCode(
        version=None,
        error_correction=ec_level,
        box_size=box_size,
        border=border,
    )
    qr.add_data(content)
    qr.make(fit=True)

    img = qr.make_image(
        image_factory=StyledPilImage,
        module_drawer=drawer_cls(),
        color_mask=SolidFillColorMask(back_color=bg, front_color=fg),
    )

    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    buffer.seek(0)
    png_bytes = buffer.getvalue()
    encoded = base64.b64encode(png_bytes).decode("ascii")
    return jsonify({
        "data_url": f"data:image/png;base64,{encoded}",
        "size": len(png_bytes),
        "modules": qr.modules_count,
    })


def _qr_record_to_data_url(rec: dict[str, Any]) -> str:
    png_blob = rec.get("png_blob") or b""
    return f"data:image/png;base64,{base64.b64encode(png_blob).decode('ascii')}"


def _qr_record_public(rec: dict[str, Any], include_image: bool = True) -> dict[str, Any]:
    out = {
        "id": rec["id"],
        "label": rec.get("label") or "",
        "content": rec.get("content") or "",
        "style": rec.get("style") or "square",
        "fg_color": rec.get("fg_color") or "#111827",
        "bg_color": rec.get("bg_color") or "#ffffff",
        "error_correction": rec.get("error_correction") or "M",
        "box_size": rec.get("box_size") or 12,
        "border": rec.get("border") or 4,
        "created_at": rec.get("created_at") or 0,
    }
    if include_image and rec.get("png_blob") is not None:
        out["data_url"] = _qr_record_to_data_url(rec)
    return out


@app.post("/api/qr-codes")
def save_qr_endpoint() -> Any:
    data = request.get_json(silent=True) or {}
    data_url = (data.get("data_url") or "").strip()
    if not data_url.startswith("data:image/png;base64,"):
        return jsonify({"error": "Missing or invalid data_url"}), 400
    try:
        png_bytes = base64.b64decode(data_url.split(",", 1)[1])
    except Exception:
        return jsonify({"error": "Invalid base64 image"}), 400
    if not png_bytes:
        return jsonify({"error": "Empty image"}), 400

    label = (data.get("label") or "").strip()[:80]
    content = (data.get("content") or "").strip()
    if not content:
        return jsonify({"error": "content is required"}), 400
    if len(content) > 2000:
        return jsonify({"error": "content too long"}), 400
    if not label:
        label = content[:40] + ("…" if len(content) > 40 else "")

    qr_id = uuid.uuid4().hex
    bookdb.save_qr_code(
        qr_id,
        label=label,
        content=content,
        style=(data.get("style") or "square")[:32],
        fg_color=(data.get("fg_color") or "#111827")[:9],
        bg_color=(data.get("bg_color") or "#ffffff")[:9],
        error_correction=(data.get("error_correction") or "M")[:2],
        box_size=int(data.get("box_size") or 12),
        border=int(data.get("border") or 4),
        png_blob=png_bytes,
    )
    rec = bookdb.get_qr_code(qr_id)
    return jsonify(_qr_record_public(rec))


@app.get("/api/qr-codes")
def list_qr_endpoint() -> Any:
    include_image = request.args.get("include_image", "1") != "0"
    items: list[dict[str, Any]] = []
    for meta in bookdb.list_qr_codes():
        if include_image:
            rec = bookdb.get_qr_code(meta["id"])
            if rec is None:
                continue
            items.append(_qr_record_public(rec, include_image=True))
        else:
            items.append(_qr_record_public(meta, include_image=False))
    return jsonify({"items": items})


@app.delete("/api/qr-codes/<qr_id>")
def delete_qr_endpoint(qr_id: str) -> Any:
    deleted = bookdb.delete_qr_code(qr_id)
    if not deleted:
        return jsonify({"error": "not found"}), 404
    return jsonify({"ok": True})


@app.post("/api/jobs/<job_id>/attach-qr")
def attach_qr_to_book(job_id: str) -> Any:
    job = _get_job(job_id)
    data = request.get_json(silent=True) or {}
    qr_id = (data.get("qr_id") or "").strip()
    heading = (data.get("heading") or "Scan this QR code").strip() or "Scan this QR code"
    caption = (data.get("caption") or "").strip()

    if not qr_id:
        return jsonify({"error": "qr_id is required"}), 400

    rec = bookdb.get_qr_code(qr_id)
    if rec is None:
        return jsonify({"error": "QR code not found"}), 404

    with job.lock:
        final_path = Path(job.final_docx or job.output_docx)
        status = job.status

    if status != "success":
        return jsonify({"error": "Book is not ready yet"}), 400
    if not final_path.exists():
        return jsonify({"error": "Final docx not found"}), 404

    # Write QR PNG to disk next to the document so it can be re-opened later.
    qr_assets_dir = OUTPUT_DIR / "qr_assets"
    qr_assets_dir.mkdir(parents=True, exist_ok=True)
    qr_png_path = qr_assets_dir / f"qr_{job_id}_{qr_id}.png"
    qr_png_path.write_bytes(rec["png_blob"])

    try:
        _append_qr_page(final_path, qr_png_path, heading=heading, caption=caption or rec.get("content", ""))
    except Exception as exc:
        return jsonify({"error": f"Failed to attach QR: {exc}"}), 500

    _append_log(job, f"Attached QR code to last page (label={rec.get('label')!r}).")

    # Regenerate Kindle / Paperback variants so the QR page is also in them.
    try:
        kindle_doc, paperback_doc = _run_kdp_formatting(job, final_path)
        with job.lock:
            job.kindle_docx = str(kindle_doc)
            job.paperback_docx = str(paperback_doc)
        _append_log(job, "Re-formatted Kindle + Paperback with QR page.")
    except Exception as exc:
        _append_log(job, f"WARNING: QR added to final doc but KDP re-format failed: {exc}")

    _sync_job_to_db(job)

    with job.lock:
        return jsonify({
            "ok": True,
            "final_docx": job.final_docx,
            "kindle_docx": job.kindle_docx,
            "paperback_docx": job.paperback_docx,
        })


def _append_qr_page(doc_path: Path, qr_png: Path, *, heading: str, caption: str) -> None:
    """Append a centred 'Scan this QR code' page to the given .docx in-place."""
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.text import WD_BREAK

    doc = Document(str(doc_path))

    # Page break so the QR section starts on its own page.
    page_break_para = doc.add_paragraph()
    page_break_para.add_run().add_break(WD_BREAK.PAGE)

    # Heading: "Scan this QR code"
    h_para = doc.add_paragraph()
    h_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h_run = h_para.add_run(heading)
    h_run.bold = True
    h_run.font.size = Pt(22)

    # Spacer
    doc.add_paragraph()

    # QR image, centred
    img_para = doc.add_paragraph()
    img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_para.add_run().add_picture(str(qr_png), width=Inches(3.0))

    # Caption (URL / content) — small and muted
    if caption:
        c_para = doc.add_paragraph()
        c_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c_run = c_para.add_run(caption)
        c_run.font.size = Pt(10)

    doc.save(str(doc_path))


def _find_bonus_paragraph(doc: Document) -> Paragraph | None:
    """Return the paragraph that contains the FREE BONUS page content."""
    for para in doc.paragraphs:
        raw_text = para.text or ""
        text = " ".join(raw_text.split()).upper()
        if not text:
            continue
        if "FREE BONUS" in text and "GET OUR NEXT BOOK" in text and "FOR FREE" in text:
            return para

    for para in doc.paragraphs:
        text = " ".join((para.text or "").split()).upper()
        if text == "FREE BONUS":
            return para

    return None


def _insert_paperback_bonus_page(doc_path: Path, qr_png: Path) -> None:
    """Insert a QR code into the existing FREE BONUS page in the paperback .docx."""
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    doc = Document(str(doc_path))
    bonus_para = _find_bonus_paragraph(doc)
    if bonus_para is None:
        raise ValueError("FREE BONUS page not found in paperback docx")

    # Preserve the existing top spacer (leading line breaks) so vertical centering stays intact.
    raw_text = "".join(r.text or "" for r in bonus_para.runs)
    prefix = raw_text.split("FREE BONUS", 1)[0] if "FREE BONUS" in raw_text else ""
    leading_breaks = prefix.count("\n")

    # Clear existing runs but keep paragraph properties.
    for child in list(bonus_para._p):
        if child.tag == qn("w:pPr"):
            continue
        bonus_para._p.remove(child)

    bonus_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if leading_breaks:
        spacer = bonus_para.add_run("\n" * leading_breaks)
        spacer.font.size = Pt(1)

    title_run = bonus_para.add_run("FREE BONUS")
    title_run.font.size = Pt(36)
    bonus_para.add_run("\n\n")

    qr_run = bonus_para.add_run()
    qr_run.add_picture(str(qr_png), width=Inches(3.0))

    bonus_para.add_run("\n\n")
    line2 = bonus_para.add_run("GET OUR NEXT BOOK")
    line2.font.size = Pt(20)
    bonus_para.add_run("\n")
    line3 = bonus_para.add_run("FOR FREE")
    line3.font.size = Pt(20)

    doc.save(str(doc_path))


def _insert_kindle_bonus_page(doc_path: Path, landing_url: str) -> None:
    """Insert a clickable link into the existing FREE BONUS page in the Kindle .docx."""
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document(str(doc_path))
    bonus_para = _find_bonus_paragraph(doc)
    if bonus_para is None:
        raise ValueError("FREE BONUS page not found in Kindle docx")

    bonus_para.add_run().add_break(WD_BREAK.LINE)
    bonus_para.add_run().add_break(WD_BREAK.LINE)

    # Build a w:hyperlink element with the URL as an external relationship.
    part = doc.part
    r_id = part.relate_to(
        landing_url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    # Blue underlined style for hyperlink
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "48")  # half-points: 24pt = 48
    rPr.append(sz)

    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), "48")
    rPr.append(szCs)

    new_run.append(rPr)
    text_el = OxmlElement("w:t")
    text_el.text = "Just Click Here!"
    new_run.append(text_el)
    hyperlink.append(new_run)

    bonus_para._p.append(hyperlink)

    doc.save(str(doc_path))


def _generate_qr_png(content: str, output_path: Path) -> None:
    """Generate a simple black-on-white QR code PNG at ``output_path``."""
    qr = qrcode.QRCode(
        version=None,
        error_correction=ERROR_CORRECT_H,
        box_size=12,
        border=4,
    )
    qr.add_data(content)
    qr.make(fit=True)

    img = qr.make_image(
        image_factory=StyledPilImage,
        module_drawer=SquareModuleDrawer(),
        color_mask=SolidFillColorMask(
            back_color=(255, 255, 255),
            front_color=(17, 24, 39),
        ),
    )
    img.save(str(output_path), format="PNG")


def _do_landing_page_qr(job: Job, page_title: str) -> None:
    """Core logic: create WP landing page, QR code, update both docx files.

    This is designed to be called inline (e.g. at the end of _run_generation)
    or from a standalone background task.  It logs progress but does NOT set
    job status or sync to DB — callers handle that.
    """
    _append_log(job, f"Starting landing page + QR workflow for '{page_title}'...")

    # 1. Create the WordPress landing page
    _append_log(job, "Creating WordPress landing page...")
    page_info = wp_landing_page.create_landing_page(
        title=page_title,
        callback=lambda step, msg: _append_log(job, msg),
    )
    landing_url = page_info["url"]
    _append_log(job, f"Landing page published: {landing_url}")

    # 2. Generate QR code for the landing page URL
    _append_log(job, "Generating QR code for landing page URL...")
    qr_assets_dir = OUTPUT_DIR / "qr_assets"
    qr_assets_dir.mkdir(parents=True, exist_ok=True)
    qr_png_path = qr_assets_dir / f"qr_landing_{job.id}.png"
    _generate_qr_png(landing_url, qr_png_path)
    _append_log(job, f"QR code saved: {qr_png_path.name}")

    # 3. Insert FREE BONUS page into paperback docx
    with job.lock:
        paperback_path = Path(job.paperback_docx) if job.paperback_docx else None

    if paperback_path and paperback_path.exists():
        _append_log(job, "Inserting QR bonus page into paperback docx...")
        _insert_paperback_bonus_page(paperback_path, qr_png_path)
        _append_log(job, "Paperback FREE BONUS page added with QR code.")
    else:
        _append_log(job, "WARNING: No paperback docx found — skipping QR insertion.")

    # 4. Insert link bonus page into kindle docx
    with job.lock:
        kindle_path = Path(job.kindle_docx) if job.kindle_docx else None

    if kindle_path and kindle_path.exists():
        _append_log(job, "Inserting link bonus page into Kindle docx...")
        _insert_kindle_bonus_page(kindle_path, landing_url)
        _append_log(job, "Kindle FREE BONUS page added with landing page link.")
    else:
        _append_log(job, "WARNING: No Kindle docx found — skipping link insertion.")

    _append_log(job, "Landing page + QR workflow complete!")
    _append_log(job, f"Landing page URL: {landing_url}")


def _run_landing_page_qr(job_id: str, page_title: str) -> None:
    """Background task wrapper for _do_landing_page_qr (manual QR button)."""
    job = _get_job(job_id)

    _set_status(job, "running", action="creating_landing_page", error="")

    try:
        _do_landing_page_qr(job, page_title)
        _set_status(job, "success", action="", error="")
        _sync_job_to_db(job)

    except Exception as exc:
        _append_log(job, f"ERROR: {exc}")
        _set_status(job, "error", action="", error=str(exc))
        _sync_job_to_db(job)


@app.post("/api/jobs/<job_id>/create-landing-qr")
def create_landing_qr(job_id: str) -> Any:
    """Create a WordPress landing page, generate a QR code for it, and insert
    the QR code into the paperback docx and a link into the Kindle docx."""
    job = _get_job(job_id)
    payload = request.get_json(silent=True) or {}
    page_title = (payload.get("page_title") or "").strip()

    if not page_title:
        return jsonify({"error": "page_title is required"}), 400
    if len(page_title) > 200:
        return jsonify({"error": "page_title too long (max 200 chars)"}), 400

    with job.lock:
        if job.status == "running":
            return jsonify({"error": "Job is busy. Wait for current task to finish."}), 409

    t = threading.Thread(
        target=_run_landing_page_qr,
        args=(job_id, page_title),
        daemon=True,
    )
    t.start()
    return jsonify({"ok": True, "page_title": page_title})


@app.post("/api/detect-pre-written")
def detect_pre_written_preview() -> Any:
    """Preview whether an uploaded/selected .docx looks already-written, so the
    upload form can pre-fill the 'already written' toggle. Does not create a
    job; the file (if uploaded here) is inspected and discarded."""
    upload = request.files.get("layout_file")
    layout_path_raw = (request.form.get("layout_path") or "").strip()

    tmp_path: Path | None = None
    try:
        if upload and upload.filename:
            if not upload.filename.lower().endswith(".docx"):
                return jsonify({"error": "File must be .docx"}), 400
            tmp_path = UPLOAD_DIR / f"_detect_{uuid.uuid4().hex[:8]}.docx"
            upload.save(str(tmp_path))
            doc_path = tmp_path
        elif layout_path_raw:
            doc_path = Path(layout_path_raw).expanduser().resolve()
            if not doc_path.exists() or doc_path.suffix.lower() != ".docx":
                return jsonify({"error": "Path must point to an existing .docx"}), 400
        else:
            return jsonify({"error": "Provide layout_file or layout_path"}), 400

        pre_written = _detect_pre_written(doc_path)
        try:
            word_count = sum(len((p.text or "").split()) for p in Document(str(doc_path)).paragraphs)
        except Exception:
            word_count = 0
        return jsonify({"pre_written": pre_written, "word_count": word_count})
    finally:
        if tmp_path is not None:
            tmp_path.unlink(missing_ok=True)


def _build_cfg_from_form(form: Any) -> tuple[dict[str, Any] | None, str | None]:
    """Build a job config dict from the submitted form, applying the same
    validation and model side effect create_job() has always done.

    Returns (cfg, None) on success or (None, error_message) on a validation
    failure. Shared by the single-book endpoint and the batch endpoint so a
    book built either way is configured identically. The config here is
    per-request (agent/model/images/etc.), not per-file — a batch applies one
    config to every selected file, which is why this reads the form once.
    """
    dashboard_mode = _normalize_dashboard_mode(form.get("dashboard_mode"))

    estimated_pages = _parse_int(form.get("estimated_pages"), 0)
    if dashboard_mode == DASHBOARD_MODE_LONG and estimated_pages < LONG_BOOK_MIN_PAGES:
        return None, f"Long Book dashboard requires estimated_pages >= {LONG_BOOK_MIN_PAGES}"

    submitted_agent = (form.get("agent") or "").strip()
    submitted_model = (form.get("openclaw_model") or "").strip()
    chosen_agent = submitted_agent or (
        LONG_BOOK_AGENT_ID if dashboard_mode == DASHBOARD_MODE_LONG else DEFAULTS["agent"]
    )
    chosen_model = submitted_model or (
        LONG_BOOK_MODEL_KEY if dashboard_mode == DASHBOARD_MODE_LONG else ""
    )

    if chosen_model:
        try:
            _set_openclaw_default_model(chosen_model)
        except subprocess.TimeoutExpired:
            return None, "OpenClaw CLI timed out while setting model. Make sure the gateway is running (openclaw gateway)."
        except Exception as exc:
            return None, f"Failed to set model '{chosen_model}': {exc}"

    cfg: dict[str, Any] = {
        "agent": chosen_agent,
        "book_premise": (form.get("book_premise") or "").strip(),
        "book_voice": (form.get("book_voice") or "").strip(),
        # Checkbox: present == on. Defaults to on for new uploads.
        "auto_book_context": form.get("auto_book_context") is not None,
        # HTML checkbox: unchecked sends nothing, checked sends a value. So
        # presence of the "images" form field == toggle is on.
        "images": form.get("images") is not None,
        "image_prompt_variant": (form.get("image_prompt_variant") or DEFAULTS["image_prompt_variant"]).strip() or DEFAULTS["image_prompt_variant"],
        "image_model": (form.get("image_model") or DEFAULTS["image_model"]).strip() or DEFAULTS["image_model"],
        "image_size": (form.get("image_size") or DEFAULTS["image_size"]).strip() or DEFAULTS["image_size"],
        "image_quality": (form.get("image_quality") or DEFAULTS["image_quality"]).strip() or DEFAULTS["image_quality"],
        "image_width": _parse_float(form.get("image_width"), float(DEFAULTS["image_width"])),
        "openai_api_key": (form.get("openai_api_key") or "").strip(),
        "title_placeholder": (form.get("title_placeholder") or "Book Title Placeholder").strip() or "Book Title Placeholder",
        "author_placeholder": (form.get("author_placeholder") or "Author Name").strip() or "Author Name",
        "estimated_pages": estimated_pages,
        "dashboard_mode": dashboard_mode,
        "openclaw_model": chosen_model,
    }

    if cfg["image_prompt_variant"] not in image_maker.PROMPT_VARIANTS:
        return None, "Invalid image prompt variant"

    return cfg, None


def _resolve_pre_written(input_doc: Path, pre_written_choice: str) -> bool:
    """Resolve the "already written" flag from the form choice, falling back to
    the word-count heuristic on "auto". Accepted values:
      "yes"/"true"/"1"/"on" -> force already-written
      "no"/"false"/"0"/"off" -> force generate
      "" / "auto" / anything else -> auto-detect
    """
    choice = (pre_written_choice or "auto").strip().lower()
    if choice in ("yes", "true", "1", "on"):
        return True
    if choice in ("no", "false", "0", "off"):
        return False
    return _detect_pre_written(input_doc)


def _create_job_record(input_doc: Path, cfg: dict[str, Any], pre_written: bool) -> Job:
    """Build a queued Job for input_doc, register it in JOBS, and persist it.

    Does NOT start the generation thread — the caller decides how to run it
    (immediately for a single job, or sequentially for a batch). This is the
    exact Job construction create_job() has always used, extracted so batch
    jobs are byte-for-byte the same as single ones.
    """
    job_id = uuid.uuid4().hex
    # Write in-place, same as the requested terminal command pattern.
    output_doc = input_doc

    job = Job(
        id=job_id,
        input_docx=str(input_doc),
        output_docx=str(output_doc),
        final_docx=str(output_doc),
        kindle_docx="",
        paperback_docx="",
        status="queued",
        config=cfg,
        pre_written=pre_written,
    )
    _append_log(job, "Job created.")
    if cfg.get("dashboard_mode") == DASHBOARD_MODE_LONG:
        _append_log(
            job,
            f"Long-book dashboard mode active (estimated pages: {cfg.get('estimated_pages', 0)}, agent: {cfg['agent']}, model: {cfg.get('openclaw_model') or ''}).",
        )
    if pre_written:
        _append_log(job, "Input file detected as already-written (contains full prose).")

    with JOBS_LOCK:
        JOBS[job_id] = job

    _sync_job_to_db(job)
    return job


def _save_uploaded_docx(upload: Any) -> tuple[Path | None, str | None]:
    """Persist an uploaded .docx to UPLOAD_DIR under a uuid-prefixed name.
    Returns (path, None) or (None, error_message)."""
    file_name = secure_filename(upload.filename)
    if not file_name.lower().endswith(".docx"):
        return None, "Uploaded file must be .docx"
    input_doc = UPLOAD_DIR / f"{uuid.uuid4().hex[:8]}_{file_name}"
    upload.save(str(input_doc))
    return input_doc, None


@app.post("/api/jobs")
def create_job() -> Any:
    upload = request.files.get("layout_file")
    layout_path_raw = (request.form.get("layout_path") or "").strip()

    if upload and upload.filename:
        input_doc, err = _save_uploaded_docx(upload)
        if err:
            return jsonify({"error": err}), 400
    elif layout_path_raw:
        input_doc = Path(layout_path_raw).expanduser().resolve()
        if not input_doc.exists() or input_doc.suffix.lower() != ".docx":
            return jsonify({"error": "layout_path must point to an existing .docx file"}), 400
    else:
        return jsonify({"error": "Provide either layout_file or layout_path"}), 400

    cfg, err = _build_cfg_from_form(request.form)
    if err:
        # Preserve the original status codes for the two special cases.
        if err.startswith("OpenClaw CLI timed out"):
            return jsonify({"error": err}), 504
        if err.startswith("Failed to set model"):
            return jsonify({"error": err}), 500
        return jsonify({"error": err}), 400

    # Confirmed in the pre-write dialog, so detection already ran — don't pay
    # for a second call that would return the same thing.
    if cfg.get("book_premise") and cfg.get("book_voice"):
        cfg["auto_book_context"] = False

    pre_written = _resolve_pre_written(input_doc, request.form.get("pre_written") or "auto")
    job = _create_job_record(input_doc, cfg, pre_written)

    t = threading.Thread(target=_run_generation, args=(job.id,), daemon=True)
    t.start()

    return jsonify({"job_id": job.id})


def _get_batch(batch_id: str) -> Batch:
    with BATCHES_LOCK:
        batch = BATCHES.get(batch_id)
    if batch is None:
        abort(404, description="Batch not found")
    return batch


def _run_batch(batch_id: str) -> None:
    """Worker thread: generate each book in the batch one at a time.

    _run_generation is blocking (it streams the writer subprocess to
    completion), so calling it in a loop gives true sequential 'one by one'
    execution: book N+1 does not start until book N reaches a terminal state.
    A book that fails or is stopped is left in its own terminal status and the
    batch continues to the next (per the chosen 'continue on failure' policy);
    _run_generation already swallows its own exceptions, but we guard the call
    anyway so one book can never take the whole batch worker down.
    """
    batch = _get_batch(batch_id)
    with batch.lock:
        job_ids = list(batch.job_ids)
        batch.status = "running"
        batch.updated_at = time.time()

    for index, job_id in enumerate(job_ids):
        with batch.lock:
            batch.current_index = index
            batch.updated_at = time.time()
        try:
            _run_generation(job_id)
        except Exception as exc:  # defensive: _run_generation handles its own
            try:
                job = _get_job(job_id)
                _append_log(job, f"ERROR: batch worker caught unexpected error: {exc}")
                _set_status(job, "error", action="", error=str(exc))
                _sync_job_to_db(job)
            except Exception:
                pass

    with batch.lock:
        batch.current_index = len(job_ids)
        batch.status = "done"
        batch.updated_at = time.time()


def _detect_identity_for_upload(
    input_doc: Path, cfg: dict[str, Any]
) -> dict[str, Any]:
    """Run the premise/voice detection for one uploaded outline.

    Preview only: no Job is created and nothing is written. Whatever the user
    already typed wins, exactly as in the writer, so the dialog shows the same
    values the run would actually use.
    """
    typed_premise = str(cfg.get("book_premise") or "").strip()
    typed_voice = str(cfg.get("book_voice") or "").strip()

    result: dict[str, Any] = {
        "premise": typed_premise,
        "voice": typed_voice,
        "source": "typed" if (typed_premise and typed_voice) else "detected",
        "detected": False,
        "error": "",
    }

    # Nothing to infer, or the user opted out.
    if (typed_premise and typed_voice) or not cfg.get("auto_book_context", True):
        result["source"] = "typed"
        return result

    try:
        doc = Document(str(input_doc))
        lines = [(p.text or "").strip() for p in doc.paragraphs if (p.text or "").strip()]
        identity = writer.infer_book_identity(
            agent_id=str(cfg.get("agent", DEFAULTS["agent"])),
            lines=lines,
            local=False,
            thinking=str(cfg.get("thinking") or DEFAULTS["thinking"]),
            timeout_s=int(cfg.get("timeout", DEFAULTS["timeout"])),
            title_hint=input_doc.stem.replace("_", " "),
        )
    except Exception as exc:
        result["error"] = str(exc)
        return result

    if not identity:
        result["error"] = "Could not determine the book type from this outline."
        return result

    result["detected"] = True
    if not typed_premise and identity.get("premise"):
        result["premise"] = identity["premise"]
    if not typed_voice and identity.get("voice"):
        result["voice"] = identity["voice"]
    return result


@app.post("/api/detect-book-identity")
def detect_book_identity() -> Any:
    """Preview what the AI thinks each uploaded outline is, before writing.

    Accepts the same multipart form as /api/jobs and /api/batches, so the
    dialog can show one section per uploaded book. Creates no jobs and spends
    one model call per file.
    """
    uploads = [u for u in request.files.getlist("layout_file") if u and u.filename]
    if not uploads:
        return jsonify({"error": "Provide at least one layout_file"}), 400

    cfg, err = _build_cfg_from_form(request.form)
    if err:
        return jsonify({"error": err}), 400

    # Reaching this endpoint IS the request to detect, so don't inherit the
    # form's checkbox convention (absent == off) — the probe form is built
    # from the config form and may not carry the checkbox at all. An explicit
    # "auto_book_context=0" still turns it off.
    cfg["auto_book_context"] = (request.form.get("auto_book_context") or "1").lower() not in {"0", "false", "no", "off"}

    pre_written_choice = request.form.get("pre_written") or "auto"

    books: list[dict[str, Any]] = []
    for upload in uploads:
        input_doc, save_err = _save_uploaded_docx(upload)
        if save_err:
            books.append({
                "filename": upload.filename,
                "premise": "", "voice": "",
                "detected": False, "source": "error",
                "error": save_err,
            })
            continue
        # An already-written book is generated with --no-text, so premise and
        # voice are never used. Skip the model call and tell the UI to leave it
        # out of the dialog. Resolved per file: a batch can mix outlines and
        # finished manuscripts, and the client's single toggle can't say which
        # is which.
        if _resolve_pre_written(input_doc, pre_written_choice):
            books.append({
                "filename": upload.filename,
                "input_path": str(input_doc),
                "premise": "", "voice": "",
                "detected": False, "source": "pre_written",
                "pre_written": True, "error": "",
            })
            continue

        info = _detect_identity_for_upload(input_doc, cfg)
        info["filename"] = upload.filename
        info["input_path"] = str(input_doc)
        info["pre_written"] = False
        books.append(info)

    return jsonify({"books": books})


@app.post("/api/batches")
def create_batch() -> Any:
    """Create a batch from multiple uploaded .docx files and generate them
    sequentially. Shares create_job()'s config building and Job construction,
    so each book in the batch is configured exactly like a single-book job."""
    uploads = [u for u in request.files.getlist("layout_file") if u and u.filename]
    if not uploads:
        return jsonify({"error": "Provide at least one layout_file"}), 400

    cfg, err = _build_cfg_from_form(request.form)
    if err:
        if err.startswith("OpenClaw CLI timed out"):
            return jsonify({"error": err}), 504
        if err.startswith("Failed to set model"):
            return jsonify({"error": err}), 500
        return jsonify({"error": err}), 400

    pre_written_choice = request.form.get("pre_written") or "auto"

    jobs: list[Job] = []
    for idx, upload in enumerate(uploads):
        input_doc, save_err = _save_uploaded_docx(upload)
        if save_err:
            return jsonify({"error": f"{upload.filename}: {save_err}"}), 400
        # Each file gets its own copy of the shared config so per-job budget
        # bookkeeping written into config doesn't bleed across books.
        job_cfg = dict(cfg)
        # Per-book premise/voice confirmed in the pre-write dialog. Sent as
        # book_premise_0, book_voice_0, book_premise_1, ... so each book in a
        # batch keeps its own identity instead of sharing one form value.
        per_premise = (request.form.get(f"book_premise_{idx}") or "").strip()
        per_voice = (request.form.get(f"book_voice_{idx}") or "").strip()
        if per_premise:
            job_cfg["book_premise"] = per_premise
        if per_voice:
            job_cfg["book_voice"] = per_voice
        # Already confirmed by the user, so don't pay for detection again.
        if per_premise and per_voice:
            job_cfg["auto_book_context"] = False
        pre_written = _resolve_pre_written(input_doc, pre_written_choice)
        job = _create_job_record(input_doc, job_cfg, pre_written)
        jobs.append(job)

    batch_id = uuid.uuid4().hex
    batch = Batch(id=batch_id, job_ids=[j.id for j in jobs])
    with BATCHES_LOCK:
        BATCHES[batch_id] = batch

    t = threading.Thread(target=_run_batch, args=(batch_id,), daemon=True)
    t.start()

    return jsonify({
        "batch_id": batch_id,
        "job_ids": [j.id for j in jobs],
        "count": len(jobs),
    })


@app.get("/api/batches/<batch_id>/status")
def batch_status(batch_id: str) -> Any:
    """Return batch progress plus a compact per-book status list so the UI can
    render 'Book N of M' and the queue beside the active book's detail panel."""
    batch = _get_batch(batch_id)
    with batch.lock:
        job_ids = list(batch.job_ids)
        current_index = batch.current_index
        batch_status_str = batch.status

    books = []
    for idx, job_id in enumerate(job_ids):
        try:
            job = _get_job(job_id)
        except Exception:
            continue
        with job.lock:
            books.append({
                "job_id": job_id,
                "index": idx,
                "title": job.custom_title.strip() or _derive_title(job.input_docx),
                "status": job.status,
                "current_action": job.current_action,
                "error": job.error,
            })

    active_job_id = (
        job_ids[current_index]
        if batch_status_str == "running" and 0 <= current_index < len(job_ids)
        else ""
    )
    return jsonify({
        "id": batch_id,
        "status": batch_status_str,
        "count": len(job_ids),
        "current_index": current_index,
        "active_job_id": active_job_id,
        "books": books,
    })


def _resolve_book_files(job_id: str) -> tuple[str, dict[str, Path]]:
    """Return (title, {kind: path}) of a book's output files, resolving from the
    in-memory Job first and falling back to the DB record for historical books.
    Mirrors download_file's resolution so a batch ZIP contains exactly what the
    per-book Download menu would offer. Only files that exist on disk are
    included, and only for a book that finished successfully — a failed or
    still-queued book has its final_docx pre-set to the raw input path at
    creation, so gating on status is what keeps that unprocessed input out of
    the archive."""
    with JOBS_LOCK:
        job = JOBS.get(job_id)
    if job is not None:
        with job.lock:
            status = job.status
            title = job.custom_title.strip() or _derive_title(job.input_docx)
            candidates = {
                "final": job.final_docx,
                "kindle": job.kindle_docx,
                "paperback": job.paperback_docx,
            }
    else:
        book = bookdb.get_book(job_id)
        if book is None:
            return job_id, {}
        status = book.get("status") or ""
        title = (book.get("title") or "").strip() or _derive_title(book.get("input_docx") or job_id)
        candidates = {
            "final": book.get("final_docx") or "",
            "kindle": book.get("kindle_docx") or "",
            "paperback": book.get("paperback_docx") or "",
        }

    if status != "success":
        return title, {}

    out: dict[str, Path] = {}
    for kind, raw in candidates.items():
        if not raw:
            continue
        p = Path(raw)
        if p.exists() and p.is_file():
            out[kind] = p
    return title, out


@app.get("/api/batches/<batch_id>/download")
def download_batch(batch_id: str) -> Any:
    """Stream a single ZIP with every finished book's output files, one folder
    per book (Final + Kindle + Paperback where they exist). Books that never
    produced any file (failed early) are skipped. 404 if nothing is ready."""
    batch = _get_batch(batch_id)
    with batch.lock:
        job_ids = list(batch.job_ids)

    buffer = io.BytesIO()
    added = 0
    used_folders: set[str] = set()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for idx, job_id in enumerate(job_ids, start=1):
            title, files = _resolve_book_files(job_id)
            if not files:
                continue
            # Numbered, filesystem-safe folder per book; keep it unique so two
            # books with the same title don't collide inside the archive.
            base = secure_filename(title) or "book"
            folder = f"{idx:02d}_{base}"
            suffix = 1
            while folder in used_folders:
                suffix += 1
                folder = f"{idx:02d}_{base}_{suffix}"
            used_folders.add(folder)

            for kind, path in files.items():
                # e.g. 01_My_Book/My_Book_final.docx
                arcname = f"{folder}/{base}_{kind}{path.suffix}"
                zf.write(str(path), arcname)
                added += 1

    if added == 0:
        abort(404, description="No finished files are available for this batch yet.")

    buffer.seek(0)
    return send_file(
        buffer,
        mimetype="application/zip",
        as_attachment=True,
        download_name=f"batch_{batch_id[:8]}.zip",
    )


@app.get("/api/jobs/<job_id>/status")
def job_status(job_id: str) -> Any:
    job = _get_job(job_id)
    with job.lock:
        return jsonify(
            {
                "id": job.id,
                "status": job.status,
                "current_action": job.current_action,
                "error": job.error,
                "input_docx": job.input_docx,
                "output_docx": job.output_docx,
                "final_docx": job.final_docx,
                "kindle_docx": job.kindle_docx,
                "paperback_docx": job.paperback_docx,
                "headings": list(job.headings),
                "listing": dict(job.listing) if job.listing else {},
                "pre_written": bool(job.pre_written),
                "book_premise": job.config.get("book_premise", "") or job.config.get("book_context", ""),
                "book_voice": job.config.get("book_voice", ""),
                "hemingway_login_required": bool(job.hemingway_login_required),
                "budget_paused": bool(job.budget_paused),
                "budget_cap_reached": bool(job.budget_paused),
                "budget_spent_usd": job.budget_spent_usd,
                "budget_limit_usd": job.budget_limit_usd,
                "created_at": job.created_at,
                "updated_at": job.updated_at,
            }
        )


@app.post("/api/jobs/<job_id>/budget/continue")
def job_budget_continue(job_id: str) -> Any:
    """Dismiss the warning while generation continues.

    Jobs paused by the previous hard-cap implementation are resumed for
    backward compatibility; already-written sections are skipped.
    """
    job = _get_job(job_id)
    with job.lock:
        legacy_paused = job.status == "budget_paused"
        warning_visible = job.budget_paused
        if job.status in {"success", "error"}:
            return jsonify({"ok": True, "resumed": False, "already_finished": True})
        if job.status == "running" and not warning_visible:
            return jsonify({"ok": True, "resumed": False, "already_dismissed": True})
        if not legacy_paused and not (job.status == "running" and warning_visible):
            return jsonify({"error": "No spending warning is active."}), 409
        current_limit = job.budget_limit_usd or float(
            job.config.get("max_spend_usd", DEFAULTS["max_spend_usd"])
        )
        job.budget_paused = False

    if legacy_paused:
        _append_log(job, "Continue selected. Resuming the previously paused generation.")
        t = threading.Thread(
            target=_resume_generation,
            args=(job_id, current_limit),
            daemon=True,
        )
        t.start()
        return jsonify({"ok": True, "resumed": True})

    _append_log(job, "Continue selected. Spending warning dismissed; generation is still running.")
    return jsonify({"ok": True, "resumed": False})


@app.post("/api/jobs/<job_id>/budget/stop")
def job_budget_stop(job_id: str) -> Any:
    """Stop the active writer and keep all sections already saved to disk."""
    job = _get_job(job_id)
    with job.lock:
        legacy_paused = job.status == "budget_paused"
        if job.status == "success":
            return jsonify({"ok": True, "stopping": False, "already_finished": True})
        if not legacy_paused and not (job.status == "running" and job.budget_paused):
            return jsonify({"error": "No spending warning is active."}), 409
        job.budget_paused = False
        job.stop_requested = not legacy_paused
        proc = job.active_process

    _append_log(job, "Stop selected. Ending generation and keeping completed sections.")

    if legacy_paused:
        t = threading.Thread(target=_finalize_legacy_stopped_job, args=(job_id,), daemon=True)
        t.start()
        return jsonify({"ok": True, "stopping": True}), 202

    if proc is not None and proc.poll() is None:
        try:
            os.killpg(proc.pid, signal.SIGTERM)
        except (ProcessLookupError, PermissionError, OSError):
            try:
                proc.terminate()
            except OSError:
                pass
    return jsonify({"ok": True, "stopping": True}), 202


@app.get("/api/jobs/<job_id>/logs")
def job_logs(job_id: str) -> Any:
    job = _get_job(job_id)
    start = _parse_int(request.args.get("from"), 0)
    with job.lock:
        start = max(0, min(start, len(job.logs)))
        lines = job.logs[start:]
        next_cursor = len(job.logs)
    return jsonify({"logs": lines, "next": next_cursor})


HEMINGWAY_LOGIN_PROC: subprocess.Popen | None = None
HEMINGWAY_LOGIN_LOCK = threading.Lock()


@app.post("/api/hemingway/login")
def hemingway_login() -> Any:
    """Open a visible browser on this machine so the user can log in to
    hemingwayapp.com. clarity_agent.py --login saves the session in the
    persistent Playwright profile, which the headless clarity scrub reuses."""
    global HEMINGWAY_LOGIN_PROC
    with HEMINGWAY_LOGIN_LOCK:
        if HEMINGWAY_LOGIN_PROC is not None and HEMINGWAY_LOGIN_PROC.poll() is None:
            return jsonify({"ok": True, "already_running": True})
        try:
            HEMINGWAY_LOGIN_PROC = subprocess.Popen(
                [str(PYTHON_BIN), str(ROOT_DIR / "clarity_agent.py"), "--login"],
                cwd=str(ROOT_DIR),
            )
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)}), 500
    return jsonify({"ok": True, "already_running": False})


@app.post("/api/jobs/<job_id>/replace-images")
def replace_images(job_id: str) -> Any:
    job = _get_job(job_id)
    payload = request.get_json(silent=True) or {}
    headings = payload.get("headings") or []
    if not isinstance(headings, list):
        return jsonify({"error": "headings must be an array"}), 400

    normalized = [str(h).strip() for h in headings if str(h).strip()]
    if not normalized:
        return jsonify({"error": "Select at least one heading"}), 400

    with job.lock:
        if job.status == "running":
            return jsonify({"error": "Job is busy. Wait for current task to finish."}), 409
        final_doc = Path(job.final_docx)

    if not final_doc.exists():
        return jsonify({"error": "Final document not found for this job"}), 400

    overrides = {
        "image_prompt_variant": str(payload.get("image_prompt_variant") or "").strip() or job.config.get("image_prompt_variant", DEFAULTS["image_prompt_variant"]),
        "image_model": str(payload.get("image_model") or "").strip() or job.config.get("image_model", DEFAULTS["image_model"]),
        "image_size": str(payload.get("image_size") or "").strip() or job.config.get("image_size", DEFAULTS["image_size"]),
        "image_quality": str(payload.get("image_quality") or "").strip() or job.config.get("image_quality", DEFAULTS["image_quality"]),
        "image_width": _parse_float(str(payload.get("image_width") or ""), float(job.config.get("image_width", DEFAULTS["image_width"]))),
        "openai_api_key": str(payload.get("openai_api_key") or "").strip() or job.config.get("openai_api_key", ""),
        "image_guidance": str(payload.get("image_guidance") or "").strip(),
    }

    if overrides["image_prompt_variant"] not in image_maker.PROMPT_VARIANTS:
        return jsonify({"error": "Invalid image prompt variant"}), 400

    t = threading.Thread(target=_run_replace_images, args=(job_id, normalized, overrides), daemon=True)
    t.start()
    return jsonify({"ok": True, "queued": len(normalized)})


@app.get("/api/jobs/<job_id>/rewritable-headings")
def rewritable_headings(job_id: str) -> Any:
    """Headings the rewrite menu can offer, including subheadings.

    Separate from job.headings (image headings, chapter-level only) because a
    chapter title has no body paragraph of its own to rewrite.
    """
    job = _get_job(job_id)
    with job.lock:
        final_doc = Path(job.final_docx or job.output_docx)

    if not final_doc.exists():
        return jsonify({"error": "Final document not found for this job"}), 400

    try:
        items = _list_rewritable_headings(final_doc)
    except Exception as exc:
        return jsonify({"error": f"Could not read document: {exc}"}), 500

    return jsonify({"headings": items})


@app.post("/api/jobs/<job_id>/rewrite-paragraphs")
def rewrite_paragraphs(job_id: str) -> Any:
    job = _get_job(job_id)
    payload = request.get_json(silent=True) or {}
    headings = payload.get("headings") or []
    if not isinstance(headings, list):
        return jsonify({"error": "headings must be an array"}), 400

    normalized = [str(h).strip() for h in headings if str(h).strip()]
    if not normalized:
        return jsonify({"error": "Select at least one heading"}), 400

    with job.lock:
        if job.status == "running":
            return jsonify({"error": "Job is busy. Wait for current task to finish."}), 409
        final_doc = Path(job.final_docx)

    if not final_doc.exists():
        return jsonify({"error": "Final document not found for this job"}), 400

    overrides = {
        "agent": str(payload.get("agent") or "").strip() or job.config.get("agent", DEFAULTS["agent"]),
        "tone": str(payload.get("tone") or "").strip() or job.config.get("tone", DEFAULTS["tone"]),
        "words": int(payload.get("words") or job.config.get("words", DEFAULTS["words"])),
        "words_max": int(payload.get("words_max") or job.config.get("words_max", DEFAULTS["words_max"])),
        "subwords": int(payload.get("subwords") or job.config.get("subwords", DEFAULTS["subwords"])),
        "subwords_max": int(payload.get("subwords_max") or job.config.get("subwords_max", DEFAULTS["subwords_max"])),
        "rewrite_guidance": str(payload.get("rewrite_guidance") or "").strip(),
        # Inherited from the job, not the rewrite dialog: without these a
        # rewritten paragraph loses the book's identity and reverts to generic
        # prose. Still overridable via the API for callers that send them.
        "book_premise": (
            str(payload.get("book_premise")).strip()
            if payload.get("book_premise") is not None
            else (job.config.get("book_premise", "") or job.config.get("book_context", ""))
        ),
        "book_voice": (
            str(payload.get("book_voice")).strip()
            if payload.get("book_voice") is not None
            else job.config.get("book_voice", "")
        ),
    }

    t = threading.Thread(target=_run_rewrite_paragraphs, args=(job_id, normalized, overrides), daemon=True)
    t.start()
    return jsonify({"ok": True, "queued": len(normalized)})


def _run_generate_listing(job_id: str, title_override: str, extra_context: str = "") -> None:
    job = _get_job(job_id)

    with job.lock:
        final_doc = Path(job.final_docx or job.output_docx)
        base_cfg = dict(job.config)

    _set_status(job, "running", action="generating_listing", error="")
    _append_log(job, "Starting publishing listing generation...")
    if extra_context.strip():
        _append_log(job, f"Using author guidance: {extra_context.strip()[:200]}")

    try:
        title = title_override or base_cfg.get("title_placeholder", "")

        def _progress(step: str, msg: str) -> None:
            _append_log(job, f"[{step}] {msg}")

        result = pub_listing_agent.generate_listing(
            docx_path=final_doc,
            title=title,
            agent_id="pub-listing-agent-1",
            timeout_s=base_cfg.get("timeout", 180),
            callback=_progress,
            extra_context=extra_context,
        )

        listing_data = {
            "title": result.title,
            "subtitles": result.subtitles,
            "description": result.description,
            "ebook_categories": result.ebook_categories,
            "paperback_categories": result.paperback_categories,
        }

        with job.lock:
            job.listing = listing_data

        _append_log(job, f"Subtitles: {len(result.subtitles)} ideas generated")
        _append_log(job, f"Description: {len(result.description.split())} words")
        _append_log(job, f"Ebook categories: {len(result.ebook_categories)}")
        _append_log(job, f"Paperback categories: {len(result.paperback_categories)}")
        _append_log(job, "Publishing listing generation complete.")
        _set_status(job, "success", action="", error="")
        _sync_job_to_db(job)
    except Exception as exc:
        _append_log(job, f"ERROR: {exc}")
        _set_status(job, "error", action="", error=str(exc))
        _sync_job_to_db(job)


@app.post("/api/jobs/<job_id>/generate-listing")
def generate_listing(job_id: str) -> Any:
    job = _get_job(job_id)
    payload = request.get_json(silent=True) or {}
    title_override = str(payload.get("title", "")).strip()
    extra_context = str(payload.get("extra_context", "")).strip()

    with job.lock:
        if job.status == "running":
            return jsonify({"error": "Job is busy. Wait for current task to finish."}), 409
        final_doc = Path(job.final_docx)

    if not final_doc.exists():
        return jsonify({"error": "Final document not found. Generate the book first."}), 400

    t = threading.Thread(
        target=_run_generate_listing,
        args=(job_id, title_override, extra_context),
        daemon=True,
    )
    t.start()
    return jsonify({"ok": True})


def _run_redo_listing_part(job_id: str, part: str, extra_context: str) -> None:
    """Regenerate just one part of the listing (description, subtitles, or categories)."""
    job = _get_job(job_id)

    with job.lock:
        final_doc = Path(job.final_docx or job.output_docx)
        base_cfg = dict(job.config)
        existing = dict(job.listing or {})

    _set_status(job, "running", action=f"redo_listing_{part}", error="")
    _append_log(job, f"Redoing {part} with new guidance...")
    if extra_context.strip():
        _append_log(job, f"Author guidance: {extra_context.strip()[:200]}")

    try:
        # Pull fresh outline/intro from the finished doc
        doc_title, outline, intro = pub_listing_agent._extract_outline_and_intro(final_doc)
        title = existing.get("title") or doc_title or base_cfg.get("title_placeholder", "") or "Untitled Book"
        timeout_s = base_cfg.get("timeout", 180)
        agent_id = "pub-listing-agent-1"

        updated = dict(existing)
        updated["title"] = title

        if part == "description":
            description, _raw = pub_listing_agent.generate_description(
                title, outline, intro, agent_id, timeout_s, extra_context=extra_context
            )
            updated["description"] = description
            _append_log(job, f"New description: {len(description.split())} words")
        elif part == "subtitles":
            subtitles, _raw = pub_listing_agent.generate_subtitles(
                title, outline, intro, agent_id, timeout_s, extra_context=extra_context
            )
            updated["subtitles"] = subtitles
            _append_log(job, f"New subtitles: {len(subtitles)} ideas")
        elif part == "categories":
            ebook_cats, pb_cats, _raw = pub_listing_agent.select_categories(
                title, outline, intro, agent_id, timeout_s, extra_context=extra_context
            )
            updated["ebook_categories"] = ebook_cats
            updated["paperback_categories"] = pb_cats
            _append_log(job, f"New categories: {len(ebook_cats)} ebook, {len(pb_cats)} paperback")
        else:
            raise ValueError(f"Unknown listing part: {part}")

        with job.lock:
            job.listing = updated

        _append_log(job, f"Redo of {part} complete.")
        _set_status(job, "success", action="", error="")
        _sync_job_to_db(job)
    except Exception as exc:
        _append_log(job, f"ERROR: {exc}")
        _set_status(job, "error", action="", error=str(exc))
        _sync_job_to_db(job)


@app.post("/api/jobs/<job_id>/redo-listing")
def redo_listing(job_id: str) -> Any:
    job = _get_job(job_id)
    payload = request.get_json(silent=True) or {}
    part = str(payload.get("part", "")).strip().lower()
    extra_context = str(payload.get("extra_context", "")).strip()

    if part not in {"description", "subtitles", "categories"}:
        return jsonify({"error": "part must be one of: description, subtitles, categories"}), 400

    with job.lock:
        if job.status == "running":
            return jsonify({"error": "Job is busy. Wait for current task to finish."}), 409
        final_doc = Path(job.final_docx)

    if not final_doc.exists():
        return jsonify({"error": "Final document not found. Generate the book first."}), 400

    t = threading.Thread(
        target=_run_redo_listing_part,
        args=(job_id, part, extra_context),
        daemon=True,
    )
    t.start()
    return jsonify({"ok": True})


@app.get("/api/jobs/<job_id>/listing")
def get_listing(job_id: str) -> Any:
    job = _get_job(job_id)
    with job.lock:
        return jsonify(job.listing or {})


@app.get("/api/jobs/<job_id>/download/<kind>")
def download_file(job_id: str, kind: str) -> Any:
    # Try in-memory job first; fall back to DB for historical books
    mapping: dict[str, Path | None] = {}
    with JOBS_LOCK:
        job = JOBS.get(job_id)
    if job:
        with job.lock:
            mapping = {
                "input": Path(job.input_docx),
                "output": Path(job.output_docx),
                "final": Path(job.final_docx),
                "kindle": Path(job.kindle_docx) if job.kindle_docx else None,
                "paperback": Path(job.paperback_docx) if job.paperback_docx else None,
            }
    else:
        book = bookdb.get_book(job_id)
        if book is None:
            abort(404, description="Job not found")
        mapping = {
            "input": Path(book["input_docx"]) if book.get("input_docx") else None,
            "final": Path(book["final_docx"]) if book.get("final_docx") else None,
            "kindle": Path(book["kindle_docx"]) if book.get("kindle_docx") else None,
            "paperback": Path(book["paperback_docx"]) if book.get("paperback_docx") else None,
        }

    target = mapping.get(kind)
    if target is None:
        abort(404)
    if not target.exists() or not target.is_file():
        abort(404, description="File does not exist yet")
    return send_file(target, as_attachment=True)


@app.get("/api/jobs/<job_id>/raw-text")
def get_raw_text(job_id: str) -> Any:
    """Return the raw AI-generated text, before formatting and the Hemingway
    clarity scrub. Generation writes into the uploaded docx itself; the
    Hemingway output is saved separately as *_formatted_clear.docx."""
    raw_path: Path | None = None
    with JOBS_LOCK:
        job = JOBS.get(job_id)
    if job:
        with job.lock:
            raw_path = Path(job.input_docx) if job.input_docx else None
    else:
        book = bookdb.get_book(job_id)
        if book is None:
            abort(404, description="Job not found")
        raw_path = Path(book["input_docx"]) if book.get("input_docx") else None

    if raw_path is None or not raw_path.exists() or not raw_path.is_file():
        abort(404, description="Raw document not found")

    doc = Document(str(raw_path))
    paragraphs = [(p.text or "").strip() for p in doc.paragraphs if (p.text or "").strip()]
    return jsonify({"file": raw_path.name, "paragraphs": paragraphs})


# ----------------------------
# Book history
# ----------------------------

@app.get("/api/books")
def list_books() -> Any:
    """Return recent books for the history sidebar."""
    limit = _parse_int(request.args.get("limit"), 50)
    mode_raw = (request.args.get("dashboard_mode") or "").strip()
    dashboard_mode: str | None = None
    if mode_raw:
        dashboard_mode = _normalize_dashboard_mode(mode_raw, default="")
        if dashboard_mode not in {DASHBOARD_MODE_STANDARD, DASHBOARD_MODE_LONG}:
            return jsonify({"error": "dashboard_mode must be 'standard' or 'long-book'"}), 400
    return jsonify({"books": bookdb.list_books(limit=limit, dashboard_mode=dashboard_mode)})


@app.get("/api/books/<book_id>")
def get_book(book_id: str) -> Any:
    """Return full book record for re-loading a past job."""
    book = bookdb.get_book(book_id)
    if book is None:
        abort(404, description="Book not found")
    return jsonify(book)


@app.delete("/api/books/<book_id>")
def delete_book(book_id: str) -> Any:
    """Delete a book from history."""
    if bookdb.delete_book(book_id):
        return jsonify({"ok": True})
    abort(404, description="Book not found")


@app.patch("/api/books/<book_id>")
def patch_book(book_id: str) -> Any:
    """Update editable fields of a book (currently: title)."""
    payload = request.get_json(silent=True) or {}
    new_title = (payload.get("title") or "").strip()
    if not new_title:
        return jsonify({"error": "title is required"}), 400
    if len(new_title) > 200:
        return jsonify({"error": "title too long (max 200 chars)"}), 400

    book = bookdb.get_book(book_id)
    if book is None:
        abort(404, description="Book not found")

    bookdb.update_book(book_id, title=new_title)

    # If this job is still in memory, keep the custom title aligned so future
    # syncs don't clobber it with the filename-derived default.
    with JOBS_LOCK:
        job = JOBS.get(book_id)
    if job is not None:
        with job.lock:
            job.custom_title = new_title

    return jsonify({"ok": True, "title": new_title})


# ----------------------------
# OpenClaw model & agent management
# ----------------------------

_MODEL_CACHE: dict[str, Any] = {"data": None, "ts": 0.0}
_AGENT_CACHE: dict[str, Any] = {"data": None, "ts": 0.0}
_MODEL_CACHE_TTL = 300  # seconds
_AGENT_CACHE_TTL = 300


@app.get("/api/openclaw-models")
def list_openclaw_models() -> Any:
    """Return available OpenClaw models + current default."""
    now = time.time()
    refresh = request.args.get("refresh") == "1"

    # Serve from cache if fresh enough
    if not refresh and _MODEL_CACHE["data"] and (now - _MODEL_CACHE["ts"]) < _MODEL_CACHE_TTL:
        return jsonify(_MODEL_CACHE["data"])

    try:
        models_proc = subprocess.run(
            ["openclaw", "models", "list", "--json"],
            capture_output=True, text=True, timeout=60,
        )
        if models_proc.returncode != 0:
            return jsonify({"error": f"openclaw models list failed: {models_proc.stderr.strip()}"}), 500
        models_data = json.loads(models_proc.stdout)

        status_proc = subprocess.run(
            ["openclaw", "models", "status", "--json"],
            capture_output=True, text=True, timeout=60,
        )
        default_model = ""
        if status_proc.returncode == 0:
            status_data = json.loads(status_proc.stdout)
            default_model = status_data.get("defaultModel", "")

        result = {
            "models": models_data.get("models", []),
            "default": default_model,
        }
        _MODEL_CACHE["data"] = result
        _MODEL_CACHE["ts"] = now
        return jsonify(result)
    except subprocess.TimeoutExpired:
        return jsonify({"error": "OpenClaw CLI timed out. Make sure the gateway is running (openclaw gateway)."}), 504
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


@app.post("/api/openclaw-models/set")
def set_openclaw_model() -> Any:
    """Switch the default OpenClaw model."""
    payload = request.get_json(silent=True) or {}
    model_key = str(payload.get("model") or "").strip()
    if not model_key:
        return jsonify({"error": "model is required"}), 400

    try:
        _set_openclaw_default_model(model_key)
        return jsonify({"ok": True, "model": model_key})
    except subprocess.TimeoutExpired:
        return jsonify({"error": "OpenClaw CLI timed out. Make sure the gateway is running (openclaw gateway)."}), 504
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


@app.get("/api/openclaw-agents")
def list_openclaw_agents() -> Any:
    """Return available OpenClaw agents."""
    now = time.time()
    refresh = request.args.get("refresh") == "1"

    if not refresh and _AGENT_CACHE["data"] and (now - _AGENT_CACHE["ts"]) < _AGENT_CACHE_TTL:
        return jsonify(_AGENT_CACHE["data"])

    try:
        proc = subprocess.run(
            ["openclaw", "agents", "list", "--json"],
            capture_output=True, text=True, timeout=60,
        )
        if proc.returncode != 0:
            return jsonify({"error": f"openclaw agents list failed: {proc.stderr.strip()}"}), 500
        agents = json.loads(proc.stdout)
        result = {"agents": agents}
        _AGENT_CACHE["data"] = result
        _AGENT_CACHE["ts"] = now
        return jsonify(result)
    except subprocess.TimeoutExpired:
        return jsonify({"error": "OpenClaw CLI timed out. Make sure the gateway is running (openclaw gateway)."}), 504
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=8080, debug=False)
