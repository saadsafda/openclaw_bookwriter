"""Book viewer / editor: read a .docx into ordered editable blocks, write
formatted hand-edits back (Word-style toolbar), and serve embedded images.

Round-trips inline formatting (bold/italic/underline/strike/sup/sub/color/
highlight/font/size) as docx runs, plus paragraph properties (alignment,
bullet/number list, style/heading, indent). Images are read-only and preserved.

Works for both:
  - books        → final / kindle / paperback / input docx (from the books table)
  - publications → kindle / paperback docx (from the publications table)

Routes registered via ``register(app)``.
"""
from __future__ import annotations

import html as htmllib
import io
import re
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Pt, RGBColor
from flask import abort, jsonify, request, send_file

import db as bookdb
import openclaw_docx_writer as writer


# ---------------------------------------------------------------------------
# Formatting <-> HTML helpers
# ---------------------------------------------------------------------------

# Map docx highlight enum <-> hex (for round-tripping the highlight color).
_HL_TO_HEX = {
    WD_COLOR_INDEX.YELLOW: "#ffff00",
    WD_COLOR_INDEX.BRIGHT_GREEN: "#00ff00",
    WD_COLOR_INDEX.TURQUOISE: "#00ffff",
    WD_COLOR_INDEX.PINK: "#ff00ff",
    WD_COLOR_INDEX.RED: "#ff0000",
    WD_COLOR_INDEX.BLUE: "#0000ff",
    WD_COLOR_INDEX.TEAL: "#008080",
    WD_COLOR_INDEX.GREEN: "#008000",
    WD_COLOR_INDEX.VIOLET: "#800080",
    WD_COLOR_INDEX.DARK_RED: "#800000",
    WD_COLOR_INDEX.DARK_YELLOW: "#808000",
    WD_COLOR_INDEX.GRAY_25: "#c0c0c0",
    WD_COLOR_INDEX.GRAY_50: "#808080",
    WD_COLOR_INDEX.BLACK: "#000000",
    WD_COLOR_INDEX.WHITE: "#ffffff",
}
_HEX_TO_HL = {v: k for k, v in _HL_TO_HEX.items()}

_STYLE_NAMES = {"Normal", "Heading 1", "Heading 2", "Heading 3", "Title", "Subtitle"}


def _hex_from_rgb(rgb) -> str:
    try:
        return "#" + str(rgb)
    except Exception:
        return ""


def _parse_color(value: str):
    """Parse '#rrggbb' or 'rgb(r,g,b)' into an RGBColor, else None."""
    if not value:
        return None
    value = value.strip()
    m = re.match(r"#([0-9a-fA-F]{6})", value)
    if m:
        h = m.group(1)
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
    m = re.match(r"rgb\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)", value)
    if m:
        return RGBColor(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def _nearest_highlight(hex_color: str):
    """Map an arbitrary hex to the nearest WD_COLOR_INDEX highlight."""
    rgb = _parse_color(hex_color)
    if rgb is None:
        return None
    if hex_color.lower() in _HEX_TO_HL:
        return _HEX_TO_HL[hex_color.lower()]
    tr, tg, tb = rgb[0], rgb[1], rgb[2]
    best, best_d = None, 1 << 30
    for hl, hx in _HL_TO_HEX.items():
        c = _parse_color(hx)
        d = (c[0] - tr) ** 2 + (c[1] - tg) ** 2 + (c[2] - tb) ** 2
        if d < best_d:
            best_d, best = d, hl
    return best


def _run_to_html(run) -> str:
    """Serialize a docx run's text + formatting to an HTML span."""
    text = htmllib.escape(run.text or "")
    if not text:
        return ""
    text = text.replace("\n", "<br>")
    styles: list[str] = []
    f = run.font
    if run.bold:
        styles.append("font-weight:bold")
    if run.italic:
        styles.append("font-style:italic")
    decos = []
    if run.underline:
        decos.append("underline")
    if f.strike:
        decos.append("line-through")
    if decos:
        styles.append("text-decoration:" + " ".join(decos))
    if f.color is not None and f.color.rgb is not None:
        styles.append("color:" + _hex_from_rgb(f.color.rgb))
    if f.name:
        styles.append(f"font-family:'{f.name}'")
    if f.size is not None:
        try:
            styles.append(f"font-size:{int(f.size.pt)}pt")
        except Exception:
            pass
    if f.highlight_color is not None and f.highlight_color in _HL_TO_HEX:
        styles.append("background-color:" + _HL_TO_HEX[f.highlight_color])

    inner = text
    if f.superscript:
        inner = f"<sup>{inner}</sup>"
    if f.subscript:
        inner = f"<sub>{inner}</sub>"
    if styles:
        return f'<span style="{";".join(styles)}">{inner}</span>'
    return inner


def _paragraph_to_html(p) -> str:
    html = "".join(_run_to_html(r) for r in p.runs)
    return html or htmllib.escape(p.text or "")


class _RunHTMLParser(HTMLParser):
    """Parse contenteditable HTML for ONE paragraph into a list of run specs.

    Each run spec: {text, bold, italic, underline, strike, sup, sub,
    color, font, size_pt, highlight}.
    """

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.runs: list[dict[str, Any]] = []
        self._stack: list[dict[str, Any]] = []

    def _effective(self) -> dict[str, Any]:
        eff: dict[str, Any] = {}
        for layer in self._stack:
            eff.update({k: v for k, v in layer.items() if v is not None})
        return eff

    def _style_from_attrs(self, attrs) -> dict[str, Any]:
        d: dict[str, Any] = {}
        ad = dict(attrs)
        style = ad.get("style", "") or ""
        decls = dict(
            (kv.split(":", 1)[0].strip().lower(), kv.split(":", 1)[1].strip())
            for kv in style.split(";") if ":" in kv
        )
        fw = decls.get("font-weight", "")
        if fw in ("bold", "bolder") or (fw.isdigit() and int(fw) >= 600):
            d["bold"] = True
        if decls.get("font-style", "") == "italic":
            d["italic"] = True
        td = decls.get("text-decoration", "") + " " + decls.get("text-decoration-line", "")
        if "underline" in td:
            d["underline"] = True
        if "line-through" in td:
            d["strike"] = True
        if "color" in decls:
            d["color"] = decls["color"]
        bg = decls.get("background-color") or decls.get("background")
        if bg:
            d["highlight"] = bg
        if "font-family" in decls:
            d["font"] = decls["font-family"].split(",")[0].strip().strip("'\"")
        if "font-size" in decls:
            fs = decls["font-size"]
            mm = re.match(r"([\d.]+)\s*(pt|px)?", fs)
            if mm:
                val = float(mm.group(1))
                unit = mm.group(2) or "px"
                d["size_pt"] = val if unit == "pt" else round(val * 0.75, 1)
        # <font color="..."> legacy
        if ad.get("color"):
            d["color"] = ad["color"]
        return d

    def handle_starttag(self, tag, attrs):
        layer: dict[str, Any] = {}
        if tag in ("b", "strong"):
            layer["bold"] = True
        elif tag in ("i", "em"):
            layer["italic"] = True
        elif tag == "u":
            layer["underline"] = True
        elif tag in ("s", "strike", "del"):
            layer["strike"] = True
        elif tag == "sup":
            layer["sup"] = True
        elif tag == "sub":
            layer["sub"] = True
        elif tag == "br":
            eff = self._effective()
            self.runs.append({**eff, "text": "\n"})
            return
        layer.update(self._style_from_attrs(attrs))
        self._stack.append(layer)

    def handle_startendtag(self, tag, attrs):
        if tag == "br":
            eff = self._effective()
            self.runs.append({**eff, "text": "\n"})

    def handle_endtag(self, tag):
        if tag in ("b", "strong", "i", "em", "u", "s", "strike", "del",
                   "sup", "sub", "span", "font", "a", "div", "p"):
            if self._stack:
                self._stack.pop()

    def handle_data(self, data):
        if data == "":
            return
        eff = self._effective()
        self.runs.append({**eff, "text": data})


def _apply_run_spec(run, spec: dict[str, Any]) -> None:
    if spec.get("bold"):
        run.bold = True
    if spec.get("italic"):
        run.italic = True
    if spec.get("underline"):
        run.underline = True
    if spec.get("strike"):
        run.font.strike = True
    if spec.get("sup"):
        run.font.superscript = True
    if spec.get("sub"):
        run.font.subscript = True
    if spec.get("color"):
        c = _parse_color(spec["color"])
        if c is not None:
            run.font.color.rgb = c
    if spec.get("font"):
        run.font.name = spec["font"]
    if spec.get("size_pt"):
        try:
            run.font.size = Pt(float(spec["size_pt"]))
        except Exception:
            pass
    if spec.get("highlight"):
        hl = _nearest_highlight(spec["highlight"])
        if hl is not None:
            run.font.highlight_color = hl


def _html_to_runs(paragraph, html: str) -> None:
    """Clear a paragraph's content and rebuild runs from contenteditable HTML."""
    # Remove existing runs but keep paragraph properties (pPr).
    for r in list(paragraph.runs):
        r._element.getparent().remove(r._element)
    parser = _RunHTMLParser()
    parser.feed(html or "")
    parser.close()
    specs = [s for s in parser.runs if s.get("text", "") != ""]
    if not specs:
        # Ensure the paragraph isn't left totally empty of runs.
        paragraph.add_run("")
        return
    for spec in specs:
        text = spec["text"]
        if text == "\n":
            # line break inside the paragraph
            if paragraph.runs:
                paragraph.runs[-1].add_break()
            continue
        run = paragraph.add_run(text)
        _apply_run_spec(run, spec)


# ---------------------------------------------------------------------------
# Block reading / writing
# ---------------------------------------------------------------------------

def _block_type(p) -> str:
    """Classify a paragraph into heading / subheading / image / body / empty."""
    if writer.paragraph_has_image(p):
        return "image"
    if writer.is_heading_paragraph(p):
        return "heading"
    if writer.is_subheading_paragraph(p):
        return "subheading"
    text = (p.text or "").strip()
    if not text:
        return "empty"
    return "body"


def _style_label(p) -> str:
    """Normalize a paragraph's style name into one the toolbar understands."""
    try:
        name = (p.style.name or "Normal").strip()
    except Exception:
        return "Normal"
    if name in _STYLE_NAMES:
        return name
    if name in ("List Bullet", "List Number"):
        return "Normal"
    return name if name.startswith("Heading") else "Normal"


def _list_kind(p) -> str:
    try:
        name = (p.style.name or "").strip()
    except Exception:
        return ""
    if name == "List Bullet":
        return "bullet"
    if name == "List Number":
        return "number"
    return ""


def _align_label(p) -> str:
    a = p.alignment
    return {
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
        WD_ALIGN_PARAGRAPH.LEFT: "left",
    }.get(a, "")


def read_blocks(path: Path) -> list[dict[str, Any]]:
    """Return the document as an ordered list of blocks.

    Editable blocks carry rich `html` plus paragraph props (align/style/list/
    indent). Images are read-only. Empty paragraphs are kept (by index) but
    rendered minimally.
    """
    doc = Document(str(path))
    blocks: list[dict[str, Any]] = []
    for i, p in enumerate(doc.paragraphs):
        btype = _block_type(p)
        editable = btype != "image"
        block: dict[str, Any] = {
            "index": i,
            "type": btype,
            "text": p.text or "",
            "editable": editable,
        }
        if btype == "image":
            block["has_image"] = True
        else:
            block["html"] = _paragraph_to_html(p)
            block["style"] = _style_label(p)
            block["list"] = _list_kind(p)
            block["align"] = _align_label(p)
            try:
                li = p.paragraph_format.left_indent
                block["indent"] = int(round(li.inches / 0.5)) if li else 0
            except Exception:
                block["indent"] = 0
        blocks.append(block)
    return blocks


def _apply_paragraph_props(doc, p, edit: dict[str, Any]) -> None:
    """Apply alignment / list / style / indent to a paragraph."""
    # Paragraph style (heading / normal / list). List styles win if a list is set.
    list_kind = edit.get("list") or ""
    style = edit.get("style") or ""
    target_style = None
    if list_kind == "bullet":
        target_style = "List Bullet"
    elif list_kind == "number":
        target_style = "List Number"
    elif style and style in _STYLE_NAMES:
        target_style = style
    if target_style:
        try:
            p.style = doc.styles[target_style]
        except Exception:
            pass
    elif style == "Normal":
        try:
            p.style = doc.styles["Normal"]
        except Exception:
            pass

    align = edit.get("align") or ""
    p.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }.get(align, None)

    try:
        indent = int(edit.get("indent") or 0)
        from docx.shared import Inches
        p.paragraph_format.left_indent = Inches(0.5 * indent) if indent > 0 else None
    except Exception:
        pass


def write_blocks(path: Path, edits: list[dict[str, Any]] | dict[int, str]) -> dict[str, Any]:
    """Apply edits to paragraphs by index. Returns a summary.

    Each edit dict may carry: index, html (or text), align, list, style, indent.
    Image paragraphs are never written. Indexes out of range / on images are
    reported in ``skipped``.

    A plain ``{index: text}`` mapping is still accepted for back-compat.
    """
    # Normalize to a list of edit dicts.
    norm: list[dict[str, Any]] = []
    if isinstance(edits, dict):
        for idx, text in edits.items():
            norm.append({"index": idx, "text": text})
    else:
        norm = list(edits)

    doc = Document(str(path))
    paras = doc.paragraphs
    changed = 0
    skipped: list[int] = []
    for edit in norm:
        try:
            idx = int(edit.get("index"))
        except (TypeError, ValueError):
            continue
        if idx < 0 or idx >= len(paras):
            skipped.append(idx)
            continue
        p = paras[idx]
        if _block_type(p) == "image":
            skipped.append(idx)
            continue
        # Inline content: prefer html, fall back to plain text.
        if "html" in edit and edit["html"] is not None:
            _html_to_runs(p, str(edit["html"]))
        elif "text" in edit and edit["text"] is not None:
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            p.add_run(str(edit["text"]))
        # Paragraph-level properties (only if any provided).
        if any(k in edit for k in ("align", "list", "style", "indent")):
            _apply_paragraph_props(doc, p, edit)
        changed += 1
    if changed:
        doc.save(str(path))
    return {"changed": changed, "skipped": skipped}


_TOC_TITLE = "Table of Contents"
_TOC_STYLE = "OpenclawTOC"  # sentinel style applied to every TOC paragraph


def _is_toc_para(p) -> bool:
    try:
        return (p.style.name or "") == _TOC_STYLE
    except Exception:
        return False


def collect_headings(path: Path) -> list[dict[str, Any]]:
    """Return ordered headings: {index, level, text}. level 1 = heading, 2 = subheading.

    Skips paragraphs that belong to an inserted TOC.
    """
    doc = Document(str(path))
    out: list[dict[str, Any]] = []
    for i, p in enumerate(doc.paragraphs):
        if _is_toc_para(p):
            continue
        t = (p.text or "").strip()
        if not t:
            continue
        bt = _block_type(p)
        if bt == "heading":
            out.append({"index": i, "level": 1, "text": t})
        elif bt == "subheading":
            out.append({"index": i, "level": 2, "text": t})
    return out


def _ensure_toc_style(doc):
    """Return the sentinel TOC paragraph style, creating it if needed."""
    from docx.enum.style import WD_STYLE_TYPE
    try:
        return doc.styles[_TOC_STYLE]
    except KeyError:
        st = doc.styles.add_style(_TOC_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        try:
            st.base_style = doc.styles["Normal"]
        except Exception:
            pass
        return st


def insert_or_refresh_toc(
    path: Path,
    *,
    at_index: int | None = None,
    replace_start: int | None = None,
    replace_end: int | None = None,
) -> dict[str, Any]:
    """Insert (or replace) a static Table of Contents.

    Placement:
      - replace_start/replace_end set → remove that paragraph range and put the
        TOC where it was (cursor had a selection).
      - at_index set → insert the TOC *before* that paragraph (cursor position).
      - neither → insert at the top of the document (default).

    Always single-instance: any previously-inserted TOC (tagged with the
    ``OpenclawTOC`` sentinel style) is removed first, so refresh never dupes.
    Anchors are held as element references, so removing the old TOC can't shift
    the chosen insertion point.
    """
    from docx.shared import Inches, Pt

    doc = Document(str(path))
    paras = doc.paragraphs
    if not paras:
        return {"ok": False, "error": "document is empty", "entries": 0}
    n = len(paras)

    # --- Resolve the anchor paragraph (insert BEFORE it); None => append at end.
    anchor = None
    to_remove: list = []
    if (replace_start is not None and replace_end is not None
            and 0 <= replace_start <= replace_end < n):
        to_remove = [paras[i] for i in range(replace_start, replace_end + 1)]
        anchor = paras[replace_end + 1] if replace_end + 1 < n else None
    elif at_index is not None and 0 <= at_index < n:
        anchor = paras[at_index]
    else:
        anchor = paras[0]

    # If the anchor is itself part of the old TOC (cursor sat inside it), move the
    # anchor to the next non-TOC paragraph so we don't lose it on removal.
    if anchor is not None and _is_toc_para(anchor):
        try:
            start_i = paras.index(anchor)
        except ValueError:
            start_i = 0
        anchor = next((p for p in paras[start_i + 1:] if not _is_toc_para(p)), None)

    # --- Remove the previous TOC, then the replaced range (skip TOC dupes).
    for p in [p for p in paras if _is_toc_para(p)]:
        p._element.getparent().remove(p._element)
    for p in to_remove:
        if _is_toc_para(p):
            continue
        try:
            p._element.getparent().remove(p._element)
        except Exception:
            pass

    # --- Collect headings from what remains.
    headings: list[tuple[int, str]] = []
    for p in doc.paragraphs:
        if _is_toc_para(p):
            continue
        t = (p.text or "").strip()
        if not t:
            continue
        bt = _block_type(p)
        if bt == "heading":
            headings.append((1, t))
        elif bt == "subheading":
            headings.append((2, t))

    toc_style = _ensure_toc_style(doc)

    def _mk(text: str):
        # Insert before the anchor, or append at the document end if anchor gone.
        if anchor is not None:
            return anchor.insert_paragraph_before(text)
        return doc.add_paragraph(text)

    title = _mk(_TOC_TITLE)
    title.style = toc_style
    for r in title.runs:
        r.bold = True
        r.font.size = Pt(18)
    for level, text in headings:
        entry = _mk(text)
        entry.style = toc_style
        try:
            entry.paragraph_format.left_indent = Inches(0.3 * (level - 1))
        except Exception:
            pass
    spacer = _mk("")
    spacer.style = toc_style

    doc.save(str(path))
    return {"ok": True, "entries": len(headings)}


def extract_image(path: Path, para_index: int):
    """Return (bytes, content_type) for the first image in a paragraph, or None."""
    doc = Document(str(path))
    paras = doc.paragraphs
    if para_index < 0 or para_index >= len(paras):
        return None
    p = paras[para_index]
    blips = p._p.xpath(".//a:blip")
    if not blips:
        return None
    embed = blips[0].get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    )
    if not embed:
        return None
    try:
        part = doc.part.related_parts[embed]
    except KeyError:
        return None
    content_type = getattr(part, "content_type", "image/png") or "image/png"
    return part.blob, content_type


# ---------------------------------------------------------------------------
# Path resolution
# ---------------------------------------------------------------------------

def _book_docx_path(book: dict[str, Any], which: str) -> str:
    mapping = {
        "final": book.get("final_docx"),
        "kindle": book.get("kindle_docx"),
        "paperback": book.get("paperback_docx"),
        "input": book.get("input_docx"),
    }
    # Sensible fallback chain when the requested one is empty.
    return (mapping.get(which)
            or book.get("final_docx")
            or book.get("kindle_docx")
            or book.get("input_docx")
            or "")


def _pub_docx_path(pub: dict[str, Any], which: str) -> str:
    mapping = {
        "kindle": pub.get("kindle_docx_path"),
        "paperback": pub.get("paperback_docx_path"),
    }
    return (mapping.get(which)
            or pub.get("kindle_docx_path")
            or pub.get("paperback_docx_path")
            or "")


# ---------------------------------------------------------------------------
# Flask registration
# ---------------------------------------------------------------------------

def register(app) -> None:  # noqa: ANN001

    # ---- Books -----------------------------------------------------------

    @app.get("/api/books/<book_id>/content")
    def get_book_content(book_id: str):  # noqa: ANN202
        book = bookdb.get_book(book_id)
        if not book:
            abort(404)
        which = (request.args.get("which") or "final").lower()
        path = _book_docx_path(book, which)
        if not path or not Path(path).exists():
            return jsonify({"error": "document not found on disk", "which": which}), 404
        return jsonify({
            "title": book.get("title") or "",
            "which": which,
            "blocks": read_blocks(Path(path)),
        })

    @app.post("/api/books/<book_id>/content")
    def save_book_content(book_id: str):  # noqa: ANN202
        book = bookdb.get_book(book_id)
        if not book:
            abort(404)
        which = (request.args.get("which") or "final").lower()
        path = _book_docx_path(book, which)
        if not path or not Path(path).exists():
            return jsonify({"error": "document not found on disk"}), 404
        edits = _parse_edits(request.get_json(silent=True) or {})
        if edits is None:
            return jsonify({"error": "body must be {edits: [{index, text}]}"}), 400
        result = write_blocks(Path(path), edits)
        return jsonify({"ok": True, **result})

    @app.get("/api/books/<book_id>/content/image/<int:para_index>")
    def get_book_image(book_id: str, para_index: int):  # noqa: ANN202
        book = bookdb.get_book(book_id)
        if not book:
            abort(404)
        which = (request.args.get("which") or "final").lower()
        path = _book_docx_path(book, which)
        if not path or not Path(path).exists():
            abort(404)
        return _serve_image(Path(path), para_index)

    @app.post("/api/books/<book_id>/toc")
    def book_toc(book_id: str):  # noqa: ANN202
        book = bookdb.get_book(book_id)
        if not book:
            abort(404)
        which = (request.args.get("which") or "final").lower()
        path = _book_docx_path(book, which)
        if not path or not Path(path).exists():
            return jsonify({"error": "document not found on disk"}), 404
        return jsonify(insert_or_refresh_toc(Path(path), **_toc_anchor(request)))

    # ---- Publications ----------------------------------------------------

    @app.get("/api/publications/<pub_id>/content")
    def get_pub_content(pub_id: str):  # noqa: ANN202
        pub = bookdb.get_publication(pub_id)
        if not pub:
            abort(404)
        which = (request.args.get("which") or "kindle").lower()
        path = _pub_docx_path(pub, which)
        if not path or not Path(path).exists():
            return jsonify({"error": "document not found on disk", "which": which}), 404
        return jsonify({
            "title": pub.get("title") or "",
            "which": which,
            "blocks": read_blocks(Path(path)),
        })

    @app.post("/api/publications/<pub_id>/content")
    def save_pub_content(pub_id: str):  # noqa: ANN202
        pub = bookdb.get_publication(pub_id)
        if not pub:
            abort(404)
        which = (request.args.get("which") or "kindle").lower()
        path = _pub_docx_path(pub, which)
        if not path or not Path(path).exists():
            return jsonify({"error": "document not found on disk"}), 404
        edits = _parse_edits(request.get_json(silent=True) or {})
        if edits is None:
            return jsonify({"error": "body must be {edits: [{index, text}]}"}), 400
        result = write_blocks(Path(path), edits)
        return jsonify({"ok": True, **result})

    @app.get("/api/publications/<pub_id>/content/image/<int:para_index>")
    def get_pub_image(pub_id: str, para_index: int):  # noqa: ANN202
        pub = bookdb.get_publication(pub_id)
        if not pub:
            abort(404)
        which = (request.args.get("which") or "kindle").lower()
        path = _pub_docx_path(pub, which)
        if not path or not Path(path).exists():
            abort(404)
        return _serve_image(Path(path), para_index)

    @app.post("/api/publications/<pub_id>/toc")
    def pub_toc(pub_id: str):  # noqa: ANN202
        pub = bookdb.get_publication(pub_id)
        if not pub:
            abort(404)
        which = (request.args.get("which") or "kindle").lower()
        path = _pub_docx_path(pub, which)
        if not path or not Path(path).exists():
            return jsonify({"error": "document not found on disk"}), 404
        return jsonify(insert_or_refresh_toc(Path(path), **_toc_anchor(request)))


# ---------------------------------------------------------------------------
# Shared helpers for the routes
# ---------------------------------------------------------------------------

def _parse_edits(body: dict[str, Any]) -> list[dict[str, Any]] | None:
    """Return a list of edit dicts (index + html/text + paragraph props)."""
    raw = body.get("edits")
    if not isinstance(raw, list):
        return None
    edits: list[dict[str, Any]] = []
    for item in raw:
        if not isinstance(item, dict):
            continue
        try:
            idx = int(item.get("index"))
        except (TypeError, ValueError):
            continue
        e: dict[str, Any] = {"index": idx}
        if "html" in item:
            e["html"] = str(item.get("html") or "")
        if "text" in item:
            e["text"] = str(item.get("text") or "")
        for k in ("align", "list", "style"):
            if k in item:
                e[k] = str(item.get(k) or "")
        if "indent" in item:
            try:
                e["indent"] = int(item.get("indent") or 0)
            except (TypeError, ValueError):
                e["indent"] = 0
        edits.append(e)
    return edits


def _toc_anchor(req) -> dict[str, Any]:
    """Extract optional TOC placement (cursor index / selection range) from body."""
    body = req.get_json(silent=True) or {}
    out: dict[str, Any] = {}

    def _int(v):
        try:
            return int(v)
        except (TypeError, ValueError):
            return None

    ai = _int(body.get("at_index"))
    rs = _int(body.get("replace_start"))
    re_ = _int(body.get("replace_end"))
    if rs is not None and re_ is not None:
        out["replace_start"] = rs
        out["replace_end"] = re_
    elif ai is not None:
        out["at_index"] = ai
    return out


def _serve_image(path: Path, para_index: int):
    result = extract_image(path, para_index)
    if result is None:
        abort(404)
    blob, content_type = result
    return send_file(io.BytesIO(blob), mimetype=content_type)
