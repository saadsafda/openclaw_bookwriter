"""Story book exporters — Markdown, DOCX, KDP 6x9 print files, and a
fact-checking sheet.

The DOCX built here is a plain manuscript: Heading 1 for chapters, Heading 2
for story titles, body text for prose. It is then handed to the existing
kdp_docx_formatter.build_kdp_documents() for 6x9 sizing, gutters, TOC and page
numbers — that module is untouched and shared as-is.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from kdp_docx_formatter import BODY_TEXT_STYLE, ensure_body_text_style
from print_hygiene import (
    PRINT_DPI,
    PrintHygieneError,
    sanitize_for_print,
    strip_control_chars,
)

from .engine import StoryBook


def _has_chapters(book: StoryBook) -> bool:
    """True when the book is grouped rather than a flat list of stories.

    A single untitled chapter is the flat case and its heading is omitted, so
    the book reads as one continuous run of stories.
    """
    return any(ch.title.strip() for ch in book.chapters)


# --------------------------------------------------------------------------
# Markdown
# --------------------------------------------------------------------------

def to_markdown(book: StoryBook) -> str:
    cfg = book.config
    grouped = _has_chapters(book)
    lines: list[str] = [f"# {cfg.book_title}", ""]

    if cfg.topic:
        lines += [f"*{cfg.topic}*", ""]

    for chapter in book.chapters:
        if grouped and chapter.title:
            lines += [f"## {chapter.title}", ""]
            if chapter.intro:
                lines += [chapter.intro, ""]
            if chapter.illustration_path:
                lines += [f"![{chapter.title}]({chapter.illustration_path})", ""]

        for story in chapter.stories:
            heading = "###" if grouped else "##"
            lines += [f"{heading} {story.number}. {story.title}", ""]

            if story.illustration_path:
                lines += [f"![{story.title}]({story.illustration_path})", ""]

            lines += [story.body, ""]

            if story.sidebar:
                lines += [f"**{cfg.sidebar_label}:** {story.sidebar}", ""]
            if story.closer:
                lines += [f"**{cfg.closer_label}:** {story.closer}", ""]

    return "\n".join(lines).rstrip() + "\n"


def write_markdown(book: StoryBook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(to_markdown(book), encoding="utf-8")
    return path


# --------------------------------------------------------------------------
# Fact-check sheet
#
# These books are sold as collections of true stories, so every claim needs to
# be traceable. This sheet pairs each story with the researcher's context, the
# sources the model drew on, and anything it flagged as uncertain — the working
# document for the verification pass before publication.
# --------------------------------------------------------------------------

def to_factcheck_markdown(book: StoryBook) -> str:
    lines: list[str] = [
        f"# Fact-check sheet — {book.config.book_title}",
        "",
        "One row per story. Verify every flagged claim against at least two "
        "independent sources before publication.",
        "",
    ]

    flagged = 0
    for story in book.all_stories():
        lines += [f"## {story.number}. {story.title}", ""]

        known = [
            (label, value) for label, value in (
                ("Who", story.who), ("Year", story.year), ("Where", story.where)
            ) if value
        ]
        if known:
            lines.append(
                " · ".join(f"**{label}:** {value}" for label, value in known)
            )
            lines.append("")

        if story.context:
            lines += ["**Researcher context:**", "", f"> {story.context}", ""]
        else:
            lines += [
                "**Researcher context:** _none supplied — this story was written "
                "from the model's own knowledge and needs full verification._",
                "",
            ]

        if story.sources:
            lines += [f"**Supplied sources:** {story.sources}", ""]
        if story.cited_sources:
            lines += ["**Sources the model cited (unverified):**", ""]
            lines += [f"- {s}" for s in story.cited_sources]
            lines.append("")

        if story.uncertain_claims:
            flagged += 1
            lines += ["**⚠ Claims flagged for verification:**", ""]
            lines += [f"- [ ] {c}" for c in story.uncertain_claims]
            lines.append("")

        if story.warnings:
            lines += ["**Build warnings:**", ""]
            lines += [f"- {w}" for w in story.warnings]
            lines.append("")

    lines.insert(4, f"**{flagged} of {len(book.all_stories())} stories have "
                    f"flagged claims.**\n")
    return "\n".join(lines).rstrip() + "\n"


def write_factcheck(book: StoryBook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(to_factcheck_markdown(book), encoding="utf-8")
    return path


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def _add_story_extra(doc: Document, label: str, text: str) -> None:
    """A sidebar or closer block: a bold run-in label, then the text.

    Indented and italic so it reads as set apart from the story body without
    needing a real text box, which the KDP formatter would have to reflow.
    """
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.25)
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(8)

    run = para.add_run(f"{label}: ")
    run.bold = True
    body = para.add_run(text)
    body.italic = True


def build_docx(
    book: StoryBook,
    path: Path,
    *,
    image_width_in: float = 4.5,
) -> Path:
    """Plain manuscript DOCX. Styling/sizing is left to the KDP formatter."""
    strip_control_chars(book)
    cfg = book.config
    grouped = _has_chapters(book)
    doc = Document()
    ensure_body_text_style(doc)

    # Front matter: title page.
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(cfg.book_title)
    title_run.bold = True
    title_run.font.size = Pt(28)

    if cfg.topic:
        # Short and unpunctuated, so the formatter would read it as a heading.
        sub = doc.add_paragraph(cfg.topic)
        sub.style = doc.styles[BODY_TEXT_STYLE]
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in sub.runs:
            run.italic = True

    doc.add_page_break()

    for chapter in book.chapters:
        if grouped and chapter.title:
            doc.add_heading(chapter.title, level=1)
            if chapter.intro:
                doc.add_paragraph(chapter.intro)
            if chapter.illustration_path and Path(chapter.illustration_path).exists():
                try:
                    sanitize_for_print(
                        chapter.illustration_path, PRINT_DPI, width_in=image_width_in
                    )
                except PrintHygieneError as exc:
                    book.warnings.append(
                        f"Illustration skipped for chapter — {exc}"
                    )
                else:
                    pic = doc.add_paragraph()
                    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pic.add_run().add_picture(
                        str(chapter.illustration_path), width=Inches(image_width_in)
                    )

        for story in chapter.stories:
            # Story titles sit one level below chapter headings in a grouped
            # book, and at the top level in a flat one, so the generated TOC
            # lists the stories either way.
            doc.add_heading(
                f"{story.number}. {story.title}", level=2 if grouped else 1
            )

            if story.illustration_path and Path(story.illustration_path).exists():
                try:
                    sanitize_for_print(
                        story.illustration_path, PRINT_DPI, width_in=image_width_in
                    )
                except PrintHygieneError as exc:
                    book.warnings.append(
                        f"Illustration skipped for story — {exc}"
                    )
                else:
                    pic = doc.add_paragraph()
                    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pic.add_run().add_picture(
                        str(story.illustration_path), width=Inches(image_width_in)
                    )

            for block in story.body.split("\n\n"):
                block = block.strip()
                if not block:
                    continue
                para = doc.add_paragraph(block)
                para.paragraph_format.space_after = Pt(6)

            if story.sidebar:
                _add_story_extra(doc, cfg.sidebar_label, story.sidebar)
            if story.closer:
                _add_story_extra(doc, cfg.closer_label, story.closer)

            doc.add_page_break()

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def build_kdp_files(
    book: StoryBook,
    source_docx: Path,
    out_dir: Path,
    *,
    author_placeholder: str = "Author Name",
) -> dict[str, str]:
    """Run the existing KDP formatter over the manuscript for 6x9 print."""
    from kdp_docx_formatter import build_kdp_documents

    stem = source_docx.stem
    kindle_out = out_dir / f"{stem}_kindle.docx"
    paperback_out = out_dir / f"{stem}_paperback.docx"

    kindle_path, paperback_path, estimated, inside = build_kdp_documents(
        source_docx=source_docx,
        kindle_output=kindle_out,
        paperback_output=paperback_out,
        estimated_pages=0,
        title_placeholder=book.config.book_title,
        author_placeholder=author_placeholder,
        outline_topics=set(),
    )
    return {
        "kindle": str(kindle_path),
        "paperback": str(paperback_path),
        "estimated_pages": str(estimated),
        "inside_margin_in": f"{inside:.3f}",
    }
