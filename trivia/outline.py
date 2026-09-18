"""Outline import — turn an existing outline document into chapter entries.

Typing a chapter scheme into the web form one box at a time is the slowest part
of setting up one of these books, and the form holds nothing but DOM state — so
a stray click loses the lot. This module reads the outline the operator almost
always already has and produces the chapter dicts the config expects.

Modelled on stories/outline.py, which solves the same problem for that
pipeline, but the unit here is a *chapter* rather than a story: a title, the
scope that tells the model what belongs in it, and optionally per-chapter
question and fact counts.

The parser is deliberately forgiving. Outlines are written by people, so it
accepts several heading styles and treats labelled lines ("Scope:", "Facts:")
as structured hints while folding unlabelled prose into the scope. Anything it
cannot classify still lands in the scope rather than being dropped — losing the
operator's work is far worse than an untidy scope field.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Optional

from .engine import TriviaError

# What can legitimately follow "Chapter"/"Part" as its number: digits, Roman
# numerals, or a spelled-out number. Restricting this matters — a bare `\w+`
# also matches ordinary words, so a scope line like "Chapter scope" was read as
# a heading for a chapter numbered "scope" and became a phantom chapter.
_NUMERAL = (
    r"\d{1,3}"
    r"|[ivxlcdm]{1,7}"
    r"|one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve|"
    r"thirteen|fourteen|fifteen|sixteen|seventeen|eighteen|nineteen|twenty|"
    r"thirty|forty|fifty|first|second|third|fourth|fifth|sixth|seventh|"
    r"eighth|ninth|tenth"
)

# An explicit chapter heading: "Chapter 2: Bank Jobs", "Part One — ...".
_CHAPTER = re.compile(
    r"^\s*(?:chapter|part|section|round)\s+(" + _NUMERAL + r")\b\s*[:.)—–-]?\s*(.*)$",
    re.IGNORECASE,
)

# A numbered heading: "1. Gross Body Facts", "1) ...", "3 — ...".
_NUMBERED = re.compile(
    r"^\s*(\d{1,3})\s*[.)\]:—–-]\s*(.+?)\s*$",
)

# "Scope: ...", "**Facts:** 30", "Questions: 20" — a labelled field line.
_LABELLED = re.compile(
    r"^\s*\**\s*(scope|about|covers|coverage|description|details|summary|notes|"
    r"context|questions|question count|trivia|trivia count|facts|fact count|"
    r"did you know|art|art hint|illustration|image)\s*"
    r"\**\s*[:：]\s*(.*)$",
    re.IGNORECASE,
)

# The same labels, matched anywhere in a line rather than anchored at the
# start, so several can be pulled off a single line.
_LABEL_ANYWHERE = re.compile(
    r"\**\s*\b(scope|about|covers|coverage|description|details|summary|notes|"
    r"context|questions|question count|trivia|trivia count|facts|fact count|"
    r"did you know|art|art hint|illustration|image)\s*"
    r"\**\s*[:：]",
    re.IGNORECASE,
)

# Lines that introduce the document rather than a chapter, so a leading
# paragraph of preamble is not mistaken for the first chapter's scope.
_SKIP_PREFIXES = (
    "book outline",
    "outline —",
    "outline -",
    "chapter outline",
    "table of contents",
    "contents",
)

# Headings that end the chapter list and begin production notes about the book.
_TRAILER_HEADINGS = (
    "notes for the",
    "notes on the",
    "production notes",
    "general notes",
    "tone target",
    "research",
)

# Which config field each recognised label feeds.
_SCOPE_LABELS = {
    "scope", "about", "covers", "coverage", "description",
    "details", "summary", "notes", "context",
}
_TRIVIA_LABELS = {"questions", "question count", "trivia", "trivia count"}
_FACT_LABELS = {"facts", "fact count", "did you know"}
_HINT_LABELS = {"art", "art hint", "illustration", "image"}

# A chapter title longer than this is almost certainly a run-on prose line that
# was never meant to be a heading.
MAX_TITLE_CHARS = 120


def _split_labels(line: str) -> list[tuple[str, str]]:
    """Every "label: value" pair on one line, in order.

    Counts are routinely written together — "Facts: 100, Questions: 50" — and
    reading only the first pair silently dropped the second, leaving the form
    on its default. Splitting before each label keeps a value that legitimately
    contains a comma intact, because only a real label ends a value.
    """
    matches = list(_LABEL_ANYWHERE.finditer(line))
    if not matches:
        return []
    pairs: list[tuple[str, str]] = []
    for i, match in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(line)
        value = _clean(line[match.end():end]).strip(" ,;")
        pairs.append((match.group(1).lower(), value))
    return pairs


def _has_scope(chapter: dict[str, Any], scope_parts: list[str]) -> bool:
    """Whether the open chapter has picked up any scope text yet."""
    return bool((chapter.get("chapter_scope") or "").strip()) or any(
        p.strip() for p in scope_parts
    )


def _is_trailer(line: str) -> bool:
    lowered = line.lower()
    return any(lowered.startswith(prefix) for prefix in _TRAILER_HEADINGS)


def _is_noise(line: str) -> bool:
    lowered = line.lower()
    return any(lowered.startswith(prefix) for prefix in _SKIP_PREFIXES)


def _clean(text: str) -> str:
    """Strip markdown emphasis and bullet markers a heading may carry."""
    s = (text or "").strip()
    s = re.sub(r"^[-*•‣]\s*", "", s)
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"^#{1,6}\s*", "", s)
    return s.strip()


def _as_count(value: str) -> Optional[int]:
    """First whole number in a label's value, or None.

    Tolerates "30 facts" and "about 25" — the operator is writing prose, not
    filling in a form.
    """
    match = re.search(r"\d{1,5}", value or "")
    if not match:
        return None
    try:
        n = int(match.group(0))
    except ValueError:
        return None
    return n if 0 <= n <= 10000 else None


def _read_docx(path: Path) -> list[str]:
    """Paragraph text from a .docx, including table cells.

    Chapter schemes are often laid out as a table with one row per chapter, so
    ignoring tables would silently return an empty outline for those documents.
    Heading styles are noted so a document that uses Word headings rather than
    numbering still yields chapter breaks.
    """
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - dependency is installed
        raise TriviaError("python-docx is not installed.") from exc

    try:
        doc = Document(str(path))
    except Exception as exc:
        raise TriviaError(f"Could not read that Word document: {exc}") from exc

    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text
        style = (para.style.name if para.style is not None else "") or ""
        # Mark real Word headings so an outline with no numbering still breaks
        # into chapters. The marker is stripped again by the line parser.
        if style.lower().startswith("heading") and text.strip():
            lines.append(f"\x00HEADING\x00{text}")
        else:
            lines.append(text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                lines.extend(p.text for p in cell.paragraphs)
    return lines


def parse_outline_lines(lines: list[str]) -> dict[str, Any]:
    """Group raw outline lines into chapter dicts the trivia config accepts.

    Returns {"chapters": [...]} using the same key names as the form, so the
    browser can drop the result straight into the chapter rows.
    """
    chapters: list[dict[str, Any]] = []
    current: Optional[dict[str, Any]] = None
    scope_parts: list[str] = []
    title_seen = False
    trailer_notes: list[str] = []
    in_trailer = False
    # A document-level title (the first heading before any chapter) is offered
    # back as the book title rather than becoming a chapter.
    doc_title = ""
    # The line under that title is the book's one-line description, which is
    # what the topic field wants — otherwise the operator retypes something
    # the outline already states.
    doc_topic = ""

    def _flush() -> None:
        nonlocal current, scope_parts
        if current is None:
            return
        extra = " ".join(p for p in scope_parts if p.strip()).strip()
        if extra:
            existing = current.get("chapter_scope", "")
            current["chapter_scope"] = f"{existing} {extra}".strip() if existing else extra
        current["chapter_scope"] = current.get("chapter_scope", "").strip()
        chapters.append(current)
        current = None
        scope_parts = []

    def _start(title: str) -> None:
        nonlocal current, title_seen
        _flush()
        current = {
            "chapter_number": len(chapters) + 1,
            "chapter_title": title,
            "chapter_scope": "",
            "illustration_prompt_hint": "",
        }
        title_seen = True

    for raw in lines:
        is_heading = False
        text = raw or ""
        if text.startswith("\x00HEADING\x00"):
            is_heading = True
            text = text[len("\x00HEADING\x00"):]

        line = _clean(text)
        if not line:
            continue

        # Once production notes start, no further chapter content follows.
        if _is_trailer(line):
            _flush()
            in_trailer = True
        if in_trailer:
            trailer_notes.append(line)
            continue

        # "Chapter 2: Bank Jobs" — checked before the numbered pattern so it is
        # never read as chapter-content numbering.
        chapter_match = _CHAPTER.match(line)
        if chapter_match and not _LABELLED.match(line):
            title = _clean(chapter_match.group(2)) or f"Chapter {chapter_match.group(1)}"
            _start(title)
            continue

        numbered = _NUMBERED.match(line)
        if numbered:
            candidate = _clean(numbered.group(2))
            # "1. Scope: ..." is a labelled field for the chapter above it,
            # not a new chapter.
            if candidate and not _LABELLED.match(candidate) and len(candidate) <= MAX_TITLE_CHARS:
                _start(candidate)
                continue
            line = candidate or line

        labelled = _LABELLED.match(line)
        if labelled and current is not None:
            for label, value in _split_labels(line):
                if label in _TRIVIA_LABELS:
                    n = _as_count(value)
                    if n is not None:
                        current["trivia_count"] = n
                elif label in _FACT_LABELS:
                    n = _as_count(value)
                    if n is not None:
                        current["fact_count"] = n
                elif label in _HINT_LABELS:
                    if value:
                        current["illustration_prompt_hint"] = value
                elif label in _SCOPE_LABELS and value:
                    scope_parts.append(value)
            continue

        # A Word heading with no numbering still starts a chapter. The first
        # one before any chapter content is the document title instead.
        #
        # The exception is a heading sitting directly under a chapter that has
        # no scope yet: operators style the scope line to match the title, so
        # treating it as a chapter break stole chapter 1's scope and shunted
        # every later chapter down a slot. Under an untitled-scope chapter the
        # rule "heading = title, the text under it = scope" wins.
        if is_heading and len(line) <= MAX_TITLE_CHARS:
            if not title_seen and current is None and not doc_title:
                doc_title = line
                continue
            if current is not None and not _has_scope(current, scope_parts):
                scope_parts.append(line)
                continue
            _start(line)
            continue

        # Anything else is prose: the open chapter's scope, or preamble.
        if current is not None:
            scope_parts.append(line)
        elif not title_seen and _is_noise(line):
            continue
        elif not title_seen and not doc_title and len(line) <= MAX_TITLE_CHARS:
            # A bare first line with nothing above it reads as the book title.
            doc_title = line
        elif not title_seen and not doc_topic:
            doc_topic = line

    _flush()

    # Drop anything that ended up with neither a title nor a scope.
    chapters = [
        c for c in chapters
        if (c.get("chapter_title") or "").strip() or (c.get("chapter_scope") or "").strip()
    ]

    if not chapters:
        raise TriviaError(
            "No chapters were found in that outline. Each chapter needs a "
            "heading like '1. Gross Body Facts' or 'Chapter 1: Gross Body Facts'."
        )

    # Renumber sequentially: the source may restart numbering per section.
    for i, chapter in enumerate(chapters, start=1):
        chapter["chapter_number"] = i

    return {
        "chapters": chapters,
        "chapter_count": len(chapters),
        "with_scope": sum(1 for c in chapters if c.get("chapter_scope")),
        "book_title": doc_title,
        "topic": doc_topic,
        # Returned rather than applied: only a human can say whether production
        # notes belong in the book's topic field.
        "notes": "\n".join(trailer_notes).strip(),
    }


def parse_outline_text(text: str) -> dict[str, Any]:
    return parse_outline_lines((text or "").splitlines())


def parse_outline_file(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return parse_outline_lines(_read_docx(path))
    if suffix in {".txt", ".md"}:
        return parse_outline_text(path.read_text(encoding="utf-8", errors="replace"))
    raise TriviaError("Outline must be a .docx, .txt or .md file.")
