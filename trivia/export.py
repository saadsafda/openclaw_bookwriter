"""Trivia book exporters — Markdown, DOCX, and KDP 6x9 print files.

The DOCX built here is a plain manuscript: Heading 1 for chapters, Heading 2
for sub-sections, body text for items. It is then handed to the existing
kdp_docx_formatter.build_kdp_documents() for 6x9 sizing, gutters, TOC and page
numbers — that module is untouched and shared as-is.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Optional

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from kdp_docx_formatter import (
    BODY_TEXT_STYLE,
    canonical_title_key,
    ensure_body_text_style,
)
from print_hygiene import (
    PRINT_DPI,
    PrintHygieneError,
    audit_tree,
    sanitize_for_print,
    strip_control_chars,
    xml_safe,
)

from .engine import (
    ANSWER_KEY_END_OF_BOOK,
    ANSWER_KEY_END_OF_CHAPTER,
    Chapter,
    TriviaBook,
)

LETTERS = ("A", "B", "C", "D")

TRIVIA_SECTION_TITLE = "Trivia"
FACTS_SECTION_TITLE = "Did You Know"
ANSWER_KEY_TITLE = "Answer Key"
INTRODUCTION_TITLE = "Introduction"
CONCLUSION_TITLE = "Conclusion"


# --------------------------------------------------------------------------
# Front and back matter
# --------------------------------------------------------------------------

def _default_introduction(cfg: Any) -> str:
    """Fallback for a book whose Introduction was never generated.

    Older books predate the generated front matter, and a provider refusal
    leaves the field empty, so the export still needs something to print.
    """
    return (
        f"This book collects trivia questions and surprising facts about "
        f"{cfg.topic}. Each chapter opens with a round of multiple-choice "
        f"questions, then a set of Did You Know facts. "
        + (
            "Answers for every chapter are gathered in the answer key at the "
            "back of the book."
            if cfg.answer_key_position == ANSWER_KEY_END_OF_BOOK
            else "Answers appear at the end of each chapter."
        )
    )


def _default_conclusion(cfg: Any) -> str:
    return (
        f"That is the end of the questions, but it does not have to be the end "
        f"of the subject. The best trivia leaves you curious about what else "
        f"you have not heard yet, and {cfg.topic} rewards anyone willing to "
        f"keep looking. Thank you for reading."
    )


def _front_matter_paragraphs(text: str, fallback: str) -> list[str]:
    from .engine import split_paragraphs

    return split_paragraphs(text) or split_paragraphs(fallback)


# --------------------------------------------------------------------------
# Markdown
# --------------------------------------------------------------------------

def to_markdown(book: TriviaBook) -> str:
    cfg = book.config
    lines: list[str] = [f"# {cfg.book_title}", ""]

    if cfg.topic:
        lines += [f"*A trivia and facts collection about {cfg.topic}.*", ""]

    lines += [f"## {INTRODUCTION_TITLE}", ""]
    for para in _front_matter_paragraphs(book.introduction, _default_introduction(cfg)):
        lines += [para, ""]

    for chapter in book.chapters:
        lines += [f"## Chapter {chapter.number} — {chapter.title}", ""]
        if chapter.illustration_path:
            lines += [f"![Chapter {chapter.number} illustration]({chapter.illustration_path})", ""]

        if chapter.trivia:
            lines += [f"### {TRIVIA_SECTION_TITLE}", ""]
            for i, q in enumerate(chapter.trivia, start=1):
                lines.append(f"**{i}. {q.question}**")
                lines.append("")
                for letter in LETTERS:
                    if letter in q.choices:
                        lines.append(f"{letter}. {q.choices[letter]}")
                lines.append("")

        if chapter.facts:
            lines += [f"### {FACTS_SECTION_TITLE}", ""]
            for f in chapter.facts:
                lines.append(f"- {f.fact}")
            lines.append("")

        if cfg.answer_key_position == ANSWER_KEY_END_OF_CHAPTER and chapter.trivia:
            lines += [f"### {ANSWER_KEY_TITLE} — Chapter {chapter.number}", ""]
            for i, q in enumerate(chapter.trivia, start=1):
                lines.append(f"{i}. {q.correct_answer} - {q.correct_text()}")
            lines.append("")

    if book.config.answer_key_position == ANSWER_KEY_END_OF_BOOK:
        lines += [f"## {ANSWER_KEY_TITLE}", ""]
        for chapter in book.chapters:
            if not chapter.trivia:
                continue
            lines += [f"### Chapter {chapter.number} — {chapter.title}", ""]
            for i, q in enumerate(chapter.trivia, start=1):
                lines.append(f"{i}. {q.correct_answer} - {q.correct_text()}")
            lines.append("")

    lines += [f"## {CONCLUSION_TITLE}", ""]
    for para in _front_matter_paragraphs(book.conclusion, _default_conclusion(cfg)):
        lines += [para, ""]

    return "\n".join(lines).rstrip() + "\n"


def write_markdown(book: TriviaBook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(to_markdown(book), encoding="utf-8")
    return path


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def _add_heading(doc: Document, text: str, level: int) -> None:
    doc.add_heading(xml_safe(text), level=level)


def _body_paragraph(doc: Document, text: str = ""):
    """A paragraph the KDP formatter must treat as body text, not a heading.

    Trivia lines look exactly like headings to the formatter's shape rules —
    "1. Which owl ..." reads as a numbered subheading, a short unpunctuated
    choice reads as an outline topic, and "C. Barn Owl" reads as a roman-numeral
    chapter title. Declaring the style is what keeps them body text.
    """
    para = doc.add_paragraph(xml_safe(text))
    para.style = doc.styles[BODY_TEXT_STYLE]
    return para


def _add_answer_key_block(doc: Document, chapters: list[Chapter], *, heading_level: int) -> None:
    for chapter in chapters:
        if not chapter.trivia:
            continue
        # "Chapter N — Title" is matched as a chapter opener and promoted to H1.
        _add_heading(doc, f"{chapter.title} (Chapter {chapter.number})", heading_level)
        for i, q in enumerate(chapter.trivia, start=1):
            para = _body_paragraph(doc)
            para.paragraph_format.space_after = Pt(2)
            para.add_run(f"{i}. ").bold = True
            para.add_run(f"{q.correct_answer} - {q.correct_text()}")


def build_docx(book: TriviaBook, path: Path, *, image_width_in: float = 4.5) -> Path:
    """Plain manuscript DOCX. Styling/sizing is left to the KDP formatter."""
    # One control character anywhere makes python-docx raise mid-write.
    strip_control_chars(book)
    cfg = book.config
    doc = Document()
    ensure_body_text_style(doc)

    # Front matter: title page.
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(cfg.book_title)
    title_run.bold = True
    title_run.font.size = Pt(28)

    if cfg.topic:
        # Short and unpunctuated, so the formatter would read it as a heading.
        sub = _body_paragraph(doc)
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub.add_run(f"A trivia and facts collection about {cfg.topic}")
        sub_run.italic = True

    doc.add_page_break()

    _add_heading(doc, INTRODUCTION_TITLE, 1)
    for para in _front_matter_paragraphs(book.introduction, _default_introduction(cfg)):
        doc.add_paragraph(xml_safe(para))
    doc.add_page_break()

    for chapter in book.chapters:
        _add_heading(doc, f"Chapter {chapter.number} — {chapter.title}", 1)

        if chapter.illustration_path and Path(chapter.illustration_path).exists():
            # Last line of defence before embedding: chapter art is AI
            # generated, so guarantee 300 DPI and no provenance metadata.
            try:
                sanitize_for_print(
                    chapter.illustration_path, PRINT_DPI, width_in=image_width_in
                )
            except PrintHygieneError as exc:
                book.warnings.append(f"Chapter {chapter.number} illustration skipped — {exc}")
            else:
                pic_para = doc.add_paragraph()
                pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic_para.add_run().add_picture(
                    str(chapter.illustration_path), width=Inches(image_width_in)
                )

        if chapter.trivia:
            _add_heading(doc, TRIVIA_SECTION_TITLE, 2)
            for i, q in enumerate(chapter.trivia, start=1):
                q_para = _body_paragraph(doc)
                q_para.paragraph_format.space_after = Pt(4)
                q_para.paragraph_format.keep_with_next = True
                # Bolding the whole question makes a wall of heavy text.
                q_para.add_run(f"{i}. ").bold = True
                q_para.add_run(q.question)
                for letter in LETTERS:
                    if letter not in q.choices:
                        continue
                    c_para = _body_paragraph(doc)
                    c_para.paragraph_format.left_indent = Inches(0.3)
                    c_para.paragraph_format.space_after = Pt(0)
                    c_para.paragraph_format.keep_with_next = letter != LETTERS[-1]
                    c_para.add_run(f"{letter}. {q.choices[letter]}")
                _body_paragraph(doc).paragraph_format.space_after = Pt(6)

        if chapter.facts:
            _add_heading(doc, FACTS_SECTION_TITLE, 2)
            for f in chapter.facts:
                # List Bullet + a short fact matches the formatter's topic shape.
                para = _body_paragraph(doc, f"\u2022 {f.fact}")
                para.paragraph_format.left_indent = Inches(0.25)
                para.paragraph_format.space_after = Pt(3)

        if cfg.answer_key_position == ANSWER_KEY_END_OF_CHAPTER and chapter.trivia:
            _add_heading(doc, f"{ANSWER_KEY_TITLE} — Chapter {chapter.number}", 2)
            for i, q in enumerate(chapter.trivia, start=1):
                para = _body_paragraph(doc)
                para.paragraph_format.space_after = Pt(2)
                para.add_run(f"{i}. ").bold = True
                para.add_run(f"{q.correct_answer} - {q.correct_text()}")

        doc.add_page_break()

    if cfg.answer_key_position == ANSWER_KEY_END_OF_BOOK:
        _add_heading(doc, ANSWER_KEY_TITLE, 1)
        _add_answer_key_block(doc, book.chapters, heading_level=2)
        doc.add_page_break()

    _add_heading(doc, CONCLUSION_TITLE, 1)
    for para in _front_matter_paragraphs(book.conclusion, _default_conclusion(cfg)):
        doc.add_paragraph(xml_safe(para))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def verify_print_images(book: TriviaBook, job_dir: Path) -> list[str]:
    """Confirm every image in ``job_dir`` is 300 DPI and metadata-free.

    Returns human-readable problems and records them on ``book.warnings`` so a
    bad asset surfaces in the build log instead of reaching KDP unnoticed.
    """
    problems: list[str] = []
    for path, issues in audit_tree(job_dir, PRINT_DPI).items():
        rel = path.relative_to(job_dir) if path.is_relative_to(job_dir) else path
        problems.append(f"{rel}: {'; '.join(issues)}")

    for message in problems:
        book.warnings.append(f"Print check — {message}")
    return problems


def build_kdp_files(
    book: TriviaBook,
    source_docx: Path,
    out_dir: Path,
    *,
    author_placeholder: str = "Author Name",
) -> dict[str, str]:
    """Run the existing KDP formatter over the trivia manuscript for 6x9 print."""
    from kdp_docx_formatter import build_kdp_documents

    stem = source_docx.stem
    kindle_out = out_dir / f"{stem}_kindle.docx"
    paperback_out = out_dir / f"{stem}_paperback.docx"

    # Turns on the formatter's outline guard for anything not declared body text.
    raw_topics = {
        INTRODUCTION_TITLE,
        CONCLUSION_TITLE,
        ANSWER_KEY_TITLE,
        TRIVIA_SECTION_TITLE,
        FACTS_SECTION_TITLE,
        book.config.book_title,
    }
    for chapter in book.chapters:
        raw_topics.add(f"Chapter {chapter.number} — {chapter.title}")
        raw_topics.add(chapter.title)
        raw_topics.add(f"{ANSWER_KEY_TITLE} — Chapter {chapter.number}")
        raw_topics.add(f"{chapter.title} (Chapter {chapter.number})")
    outline_topics = {canonical_title_key(t) for t in raw_topics if t}

    kindle_path, paperback_path, estimated, inside = build_kdp_documents(
        source_docx=source_docx,
        kindle_output=kindle_out,
        paperback_output=paperback_out,
        estimated_pages=0,
        title_placeholder=book.config.book_title,
        author_placeholder=author_placeholder,
        outline_topics=outline_topics,
    )
    return {
        "kindle": str(kindle_path),
        "paperback": str(paperback_path),
        "estimated_pages": str(estimated),
        "inside_margin_in": f"{inside:.3f}",
    }
