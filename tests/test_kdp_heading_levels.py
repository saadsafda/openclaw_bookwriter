"""Chapters must be Heading 1 and the outline's bullets Heading 2.

The formatter used to flatten the source outline into one undifferentiated set
of titles, so it could ask "is this line in the outline?" but never "at which
level?". Both levels then fell to text patterns, which get it wrong in both
directions: a chapter titled as a plain phrase matches no chapter pattern and
was demoted to a subheading, and a bullet opening with a number was promoted to
a chapter. The outline states the book's structure -- the formatter has to use
it.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt

import kdp_docx_formatter as K


def _outline(path: Path, rows: list[tuple[str, str]]) -> Path:
    """Write an outline whose paragraphs carry the given built-in styles."""
    doc = Document()
    for text, style in rows:
        doc.add_paragraph(text, style=style)
    doc.save(str(path))
    return path


def _styled_outline(path: Path, rows: list[tuple[str, float, bool]]) -> Path:
    """Write an outline that carries structure only in its formatting.

    This is the common real-world shape: a custom chapter style with no outline
    level, and plain body paragraphs for the bullets under it.
    """
    doc = Document()
    for text, size_pt, bold in rows:
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.font.size = Pt(size_pt)
        run.font.bold = bold
    doc.save(str(path))
    return path


def _headings(path: Path) -> list[tuple[str, str]]:
    out = []
    for para in Document(str(path)).paragraphs:
        style = (para.style.name or "")
        text = (para.text or "").strip()
        if style.lower().startswith("heading") and text and "TABLE OF CONTENTS" not in text:
            out.append((style, text))
    return out


def _build(tmp_path: Path, lines: list[str], topics) -> Path:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    src = tmp_path / "ms.docx"
    doc.save(str(src))
    paperback = tmp_path / "paperback.docx"
    K.build_kdp_documents(
        source_docx=src,
        kindle_output=tmp_path / "kindle.docx",
        paperback_output=paperback,
        estimated_pages=0,
        title_placeholder="T",
        author_placeholder="A",
        outline_topics=topics,
    )
    return paperback


def test_outline_levels_survive_loading(tmp_path):
    topics = K.load_outline_topics(str(_outline(tmp_path / "o.docx", [
        ("Finding Your Voice", "Heading 1"),
        ("Warm-up habits", "List Bullet"),
    ])))
    assert topics.level_of("FINDING YOUR VOICE") == 1
    assert topics.level_of("WARM-UP HABITS") == 2
    # Still a set, so the pipelines that pass a plain set keep working.
    assert "FINDING YOUR VOICE" in topics


def test_chapter_without_a_chapter_prefix_stays_heading_1(tmp_path):
    """A plain-phrase chapter title matches no chapter regex and was demoted."""
    topics = K.load_outline_topics(str(_outline(tmp_path / "o.docx", [
        ("Finding Your Voice", "Heading 1"),
        ("Warm-up habits", "Normal"),
        ("Working With Others", "Heading 1"),
    ])))
    built = _build(tmp_path, [
        "Finding Your Voice",
        "Warm-up habits",
        "Ordinary body prose that runs on long enough to never be a heading.",
        "Working With Others",
    ], topics)

    assert _headings(built) == [
        ("Heading 1", "Finding Your Voice"),
        ("Heading 2", "Warm-up habits"),
        ("Heading 1", "Working With Others"),
    ]


def test_a_numbered_bullet_is_not_promoted_to_a_chapter(tmp_path):
    """"1. Three warm-ups" matches the numbered-title pattern; the outline says
    it is a bullet, and the outline wins."""
    topics = K.load_outline_topics(str(_outline(tmp_path / "o.docx", [
        ("Getting Started", "Heading 1"),
        ("1. Three warm-ups", "List Bullet"),
        ("Chapter conventions explained", "List Bullet"),
    ])))
    built = _build(tmp_path, [
        "Getting Started", "1. Three warm-ups", "Chapter conventions explained",
    ], topics)

    assert _headings(built) == [
        ("Heading 1", "Getting Started"),
        ("Heading 2", "1. Three warm-ups"),
        ("Heading 2", "Chapter conventions explained"),
    ]


def test_levels_are_read_from_formatting_when_styles_say_nothing(tmp_path):
    """The real outlines use a custom style carrying no outline level, so the
    only structural signal is that chapters are set larger than the bullets."""
    topics = K.load_outline_topics(str(_styled_outline(tmp_path / "o.docx", [
        ("Welcome to the World of Acting", 18.0, True),
        ("What acting really is", 12.0, False),
        ("Why anyone can learn to act", 12.0, False),
    ])))
    assert topics.level_of("WELCOME TO THE WORLD OF ACTING") == 1
    assert topics.level_of("WHAT ACTING REALLY IS") == 2


def test_a_featureless_outline_infers_nothing(tmp_path):
    """With no signal at all, inventing a structure would be worse than falling
    back to the text heuristics."""
    topics = K.load_outline_topics(str(_outline(tmp_path / "o.docx", [
        ("Chapter 1: Beginnings", "Normal"),
        ("What it means", "Normal"),
    ])))
    assert topics.levels == {}
    assert len(topics) == 2

    built = _build(tmp_path, ["Chapter 1: Beginnings", "What it means"], topics)
    assert _headings(built) == [
        ("Heading 1", "Chapter 1: Beginnings"),
        ("Heading 2", "What it means"),
    ]


def test_kindle_and_paperback_agree_on_levels(tmp_path):
    """Both deliverables run the same detection, so the TOCs must match."""
    topics = K.load_outline_topics(str(_outline(tmp_path / "o.docx", [
        ("Finding Your Voice", "Heading 1"),
        ("Warm-up habits", "Normal"),
    ])))
    doc = Document()
    for line in ("Finding Your Voice", "Warm-up habits"):
        doc.add_paragraph(line)
    src = tmp_path / "ms.docx"
    doc.save(str(src))

    K.build_kdp_documents(
        source_docx=src,
        kindle_output=tmp_path / "kindle.docx",
        paperback_output=tmp_path / "paperback.docx",
        estimated_pages=0,
        title_placeholder="T",
        author_placeholder="A",
        outline_topics=topics,
    )
    assert _headings(tmp_path / "kindle.docx") == _headings(tmp_path / "paperback.docx")
    assert ("Heading 1", "Finding Your Voice") in _headings(tmp_path / "kindle.docx")
