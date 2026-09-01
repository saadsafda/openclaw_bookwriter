"""The Introduction and Conclusion are authored prose, not a stock paragraph.

A catalogue of books that all open with the same sentence reads as machine
output to a browsing reader, so both pieces are generated per book. The export
still has to produce something for books written before that existed, and for a
build where the provider refused the request.
"""

from __future__ import annotations

from docx import Document

from trivia import export as ex
from trivia.engine import (
    BookConfig,
    Chapter,
    ChapterConfig,
    DidYouKnowFact,
    TriviaBook,
    TriviaQuestion,
    build_conclusion_prompt,
    build_introduction_prompt,
    clean_prose_reply,
    split_paragraphs,
    word_count,
)


def _book(introduction: str = "", conclusion: str = "") -> TriviaBook:
    cfg = BookConfig(
        book_title="Bird Trivia",
        topic="birds",
        chapters=[
            ChapterConfig.from_dict({"chapter_title": "Owls", "chapter_scope": "owls"}, 1)
        ],
    )
    chapter = Chapter(number=1, title="Owls", scope="owls")
    chapter.trivia.append(
        TriviaQuestion(
            "q1", 1, "Which owl is the largest by wingspan?",
            {"A": "Great Horned", "B": "Great Gray", "C": "Barn", "D": "Snowy"},
            "B",
        )
    )
    chapter.facts.append(
        DidYouKnowFact("f1", 1, "Owl flight feathers have comb-like leading edges.")
    )
    return TriviaBook(
        config=cfg,
        chapters=[chapter],
        introduction=introduction,
        conclusion=conclusion,
    )


def _headings(path) -> list[str]:
    return [
        p.text for p in Document(str(path)).paragraphs
        if p.style.name.startswith("Heading")
    ]


def _texts(path) -> list[str]:
    return [p.text for p in Document(str(path)).paragraphs]


# -- prompts ---------------------------------------------------------------

def test_prompts_ask_for_the_full_length_and_plain_prose():
    cfg = _book().config
    for prompt in (
        build_introduction_prompt(cfg, cfg.chapters),
        build_conclusion_prompt(cfg, cfg.chapters),
    ):
        assert "300" in prompt and "500" in prompt
        # A JSON reply here would print as literal braces on the page.
        assert "JSON" not in prompt
        assert "birds" in prompt
        assert "Owls" in prompt


def test_prompts_forbid_breaking_the_fourth_wall():
    cfg = _book().config
    prompt = build_introduction_prompt(cfg, cfg.chapters)
    assert "Never mention AI" in prompt.replace("\n", " ").replace("  ", " ")


# -- reply cleaning --------------------------------------------------------

def test_clean_prose_reply_strips_fences_and_restated_headings():
    assert clean_prose_reply("```\nHello there.\n```") == "Hello there."
    assert clean_prose_reply("## Introduction\n\nHello.") == "Hello."
    assert clean_prose_reply("**Conclusion**\n\nHello.") == "Hello."
    assert clean_prose_reply("Introduction:\n\nHello.") == "Hello."


def test_split_paragraphs_handles_both_break_styles():
    assert split_paragraphs("One.\n\nTwo.") == ["One.", "Two."]
    assert split_paragraphs("One.\nTwo.") == ["One.", "Two."]
    assert split_paragraphs("   ") == []


def test_word_count_counts_hyphenated_and_contracted_words_once():
    assert word_count("A well-known fact, isn't it?") == 5


# -- export ----------------------------------------------------------------

def test_docx_prints_generated_prose_as_separate_paragraphs(tmp_path):
    book = _book("First intro para.\n\nSecond intro para.", "Closing para.")
    path = ex.build_docx(book, tmp_path / "b.docx")

    texts = _texts(path)
    assert "First intro para." in texts
    assert "Second intro para." in texts
    assert "Closing para." in texts


def test_conclusion_is_the_last_section_after_the_answer_key(tmp_path):
    book = _book("Intro.", "Outro.")
    heads = _headings(ex.build_docx(book, tmp_path / "b.docx"))

    assert heads.index("Introduction") < heads.index("Chapter 1 — Owls")
    assert heads.index("Answer Key") < heads.index("Conclusion")
    assert heads[-1] == "Conclusion"


def test_markdown_carries_both_sections(tmp_path):
    md = ex.to_markdown(_book("Intro para.", "Outro para."))

    assert "## Introduction" in md
    assert "## Conclusion" in md
    assert "Intro para." in md
    assert "Outro para." in md
    assert md.index("## Introduction") < md.index("## Chapter 1")
    assert md.index("## Conclusion") > md.index("## Chapter 1")


def test_empty_prose_falls_back_so_old_books_still_export(tmp_path):
    """A book generated before this feature has neither field set."""
    book = _book()

    md = ex.to_markdown(book)
    assert "## Introduction" in md and "## Conclusion" in md
    assert "collects trivia questions" in md

    texts = _texts(ex.build_docx(book, tmp_path / "b.docx"))
    assert any("collects trivia questions" in t for t in texts)
    assert any("does not have to be the end" in t for t in texts)


def test_both_sections_reach_the_kdp_outline_guard(tmp_path):
    """Without this they are re-shaped by the formatter's heading rules."""
    book = _book("Intro.", "Outro.")
    src = ex.build_docx(book, tmp_path / "src.docx")
    ex.build_kdp_files(book, src, tmp_path)

    heads = _headings(tmp_path / "src_paperback.docx")
    assert "Introduction" in heads
    assert "Conclusion" in heads


# -- editing ---------------------------------------------------------------

def test_front_matter_is_editable_by_hand():
    """The build writes the prose, but an operator has the last word on it."""
    from trivia import edit as editor

    book = _book(introduction="Original opening.", conclusion="Original close.")

    assert editor.apply_front_matter_edit(
        book, "introduction", {"text": "A better opening."}
    ) == "A better opening."
    assert book.introduction == "A better opening."

    # A model's restated heading is stripped on the way in, the same as during
    # generation, so pasting a reply straight from a chat window is safe.
    editor.apply_front_matter_edit(book, "conclusion", {"text": "## Conclusion\n\nA close."})
    assert book.conclusion == "A close."


def test_clearing_front_matter_restores_the_default_paragraph():
    """Blanking is the only way to undo a bad generation without a rebuild."""
    from trivia import edit as editor

    book = _book(introduction="Something wrong.")
    editor.apply_front_matter_edit(book, "introduction", {"text": "   "})
    assert book.introduction == ""
    # The export still has to put *something* before Chapter 1.
    assert "## Introduction" in ex.to_markdown(book)


def test_front_matter_edit_rejects_bad_input():
    from trivia import edit as editor
    from trivia.engine import TriviaError

    book = _book()
    for section, payload in (("chapters", {"text": "x"}), ("introduction", {})):
        try:
            editor.apply_front_matter_edit(book, section, payload)
        except TriviaError:
            continue
        raise AssertionError(f"{section}/{payload} should have been rejected")
